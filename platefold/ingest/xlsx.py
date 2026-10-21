"""Read a bounded XLSX workbook into semantic sheet sections and tables."""

from __future__ import annotations

import datetime as dt
import hashlib
import html
import json
import math
import posixpath
import re
import warnings as python_warnings
from collections import Counter
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

from markdown.extensions.toc import slugify, unique
from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, range_boundaries

from ..cancellation import CancelCheck, CancelledError, checked, never_cancelled
from ..config import Settings
from ..limits import ResourceLimitError
from ..ooxml import InvalidOOXMLError, OOXMLPackage, drawingml_theme
from . import Ingested
from .xlsx_format import format_number, format_text, general

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"

_ALLOWED_EXTERNAL_RELATIONSHIPS = frozenset(
    {"externalLink", "externalLinkPath", "hyperlink", "image", "oleObject"}
)
_IMAGE_SUFFIXES = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".svg", ".tif", ".tiff", ".webp"}
# A selector with this prefix is an exact sheet name, never an index.
SHEET_NAME_PREFIX = "name:"
_SHEET_INDEX = re.compile(r"#([1-9][0-9]*)")
_BARE_SHEET_INDEX = re.compile(r"([1-9][0-9]*)")
_FORMATS_NAMED = 3
_FORMAT_NAME_LIMIT = 60
_EXTERNAL_FORMULA = re.compile(r"\[[^\]\r\n]+\][^!\r\n]*!")


@dataclass(frozen=True)
class _Sheet:
    name: str
    state: str
    part: str
    index: int


@dataclass(frozen=True)
class _Merge:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


@dataclass(frozen=True)
class _SheetShape:
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    hidden_rows: frozenset[int]
    hidden_columns: frozenset[int]
    header: bool
    merges: tuple[_Merge, ...]
    formulas: frozenset[tuple[int, int]]
    external_formulas: frozenset[tuple[int, int]]

    @property
    def cells(self) -> int:
        if not self.max_row or not self.max_col:
            return 0
        return (self.max_row - self.min_row + 1) * (self.max_col - self.min_col + 1)


def _rels_part(part: str) -> str:
    path = PurePosixPath(part)
    return str(path.parent / "_rels" / f"{path.name}.rels")


def _resolve_part(owner: str, target: str) -> str:
    if target.startswith("/"):
        resolved = target.lstrip("/")
    else:
        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
    if resolved.startswith("../") or resolved in {"", ".", ".."}:
        raise InvalidOOXMLError(f"OOXML relationship escapes the package: {owner}")
    return resolved


def _relationships(package: OOXMLPackage, owner: str) -> dict[str, tuple[str, str, bool]]:
    part = _rels_part(owner)
    root = package.xml(part)
    if root is None:
        return {}
    result: dict[str, tuple[str, str, bool]] = {}
    for node in root.iter(f"{{{_PACKAGE_REL_NS}}}Relationship"):
        relationship_id = node.get("Id", "")
        kind = node.get("Type", "").rsplit("/", 1)[-1]
        external = node.get("TargetMode", "").casefold() == "external"
        target = "" if external else _resolve_part(owner, node.get("Target", ""))
        if not relationship_id or relationship_id in result:
            raise InvalidOOXMLError(f"Invalid duplicate OOXML relationship in {part}")
        result[relationship_id] = kind, target, external
    return result


def _sheets(package: OOXMLPackage) -> list[_Sheet]:
    workbook = package.xml("xl/workbook.xml")
    if workbook is None:
        raise InvalidOOXMLError("XLSX package has no xl/workbook.xml part")
    relationships = _relationships(package, "xl/workbook.xml")
    result: list[_Sheet] = []
    for index, node in enumerate(workbook.findall(f".//{{{_SHEET_NS}}}sheet"), 1):
        name = node.get("name", "")
        relationship_id = node.get(f"{{{_DOC_REL_NS}}}id", "")
        relation = relationships.get(relationship_id)
        if not name or relation is None:
            raise InvalidOOXMLError("XLSX workbook contains an invalid sheet declaration")
        kind, part, external = relation
        if external or kind not in {"worksheet", "chartsheet"} or not package.has(part):
            raise InvalidOOXMLError(f"XLSX workbook sheet {index} has an invalid package target")
        result.append(_Sheet(name, node.get("state", "visible"), part, index))
    if not result:
        raise InvalidOOXMLError("XLSX workbook contains no sheets")
    return result


def _resolve_selector(
    sheets: list[_Sheet], by_name: dict[str, _Sheet], selector: str
) -> _Sheet | None:
    """Resolve one selector the same way for the CLI, settings files and HTTP.

    `name:` forces an exact name; Excel forbids a colon in a sheet name, so the
    prefix cannot be part of one. `#` and a whole number from 1 is always an
    index. Any other text is an exact name, and a bare whole number that no
    sheet is named falls back to an index.
    """
    if selector.startswith(SHEET_NAME_PREFIX):
        return by_name.get(selector[len(SHEET_NAME_PREFIX) :])
    index_match = _SHEET_INDEX.fullmatch(selector)
    if index_match is None and selector in by_name:
        return by_name[selector]
    index_match = index_match or _BARE_SHEET_INDEX.fullmatch(selector)
    if index_match is None:
        return None
    index = int(index_match.group(1))
    return sheets[index - 1] if index <= len(sheets) else None


def _selected_sheets(sheets: list[_Sheet], selectors: tuple[str, ...]) -> list[_Sheet]:
    visible = [
        sheet
        for sheet in sheets
        if sheet.state == "visible" and sheet.part.startswith("xl/worksheets/")
    ]
    if not selectors:
        return visible
    by_name = {sheet.name: sheet for sheet in sheets}
    selected: list[_Sheet] = []
    for selector in selectors:
        item = _resolve_selector(sheets, by_name, selector)
        if item is None:
            hint = ""
            if selector in by_name:
                hint = f"; select the sheet named {selector!r} with {SHEET_NAME_PREFIX}{selector}"
            raise ValueError(f"Unknown XLSX sheet selection: {selector}{hint}")
        if item.state != "visible":
            raise ValueError(f"Selected XLSX sheet is hidden and cannot be exported: {selector}")
        if not item.part.startswith("xl/worksheets/"):
            raise ValueError(
                f"Selected XLSX sheet is a chart sheet and cannot be exported: {selector}"
            )
        if item not in selected:
            selected.append(item)
    return selected


def _coordinate(value: str) -> tuple[int, int]:
    try:
        column, row = coordinate_from_string(value)
        return int(row), column_index_from_string(column)
    except (TypeError, ValueError):
        raise InvalidOOXMLError(f"Invalid XLSX cell coordinate: {value!r}") from None


def _merged_ranges(
    root: Any, sheet: _Sheet, limit: int, cancel_check: CancelCheck = never_cancelled
) -> tuple[_Merge, ...]:
    merges: list[_Merge] = []
    for node in checked(root.iter(f"{{{_SHEET_NS}}}mergeCell"), cancel_check):
        if len(merges) >= limit:
            raise ResourceLimitError(
                f"XLSX sheet {sheet.index} exceeds the {limit}-merged-range limit"
            )
        try:
            min_col, min_row, max_col, max_row = range_boundaries(node.get("ref", ""))
        except (TypeError, ValueError):
            raise InvalidOOXMLError(
                f"Invalid merged-cell range in XLSX sheet {sheet.index}"
            ) from None
        if not (
            isinstance(min_row, int)
            and isinstance(max_row, int)
            and isinstance(min_col, int)
            and isinstance(max_col, int)
        ):
            raise InvalidOOXMLError(f"Invalid merged-cell range in XLSX sheet {sheet.index}")
        if (
            min_row < 1
            or max_row > 1_048_576
            or min_col < 1
            or max_col > 16_384
            or max_row < min_row
            or max_col < min_col
        ):
            raise InvalidOOXMLError(f"Invalid merged-cell range in XLSX sheet {sheet.index}")
        merges.append(_Merge(min_row, max_row, min_col, max_col))
    return tuple(merges)


def _formula_cells(
    root: Any, cancel_check: CancelCheck = never_cancelled
) -> tuple[frozenset[tuple[int, int]], frozenset[tuple[int, int]]]:
    formulas: set[tuple[int, int]] = set()
    external: set[tuple[int, int]] = set()
    shared_external: set[str] = set()
    formula_nodes: list[tuple[tuple[int, int], Any]] = []
    for cell in checked(root.iter(f"{{{_SHEET_NS}}}c"), cancel_check):
        formula = cell.find(f"{{{_SHEET_NS}}}f")
        if formula is None:
            continue
        coordinate = _coordinate(cell.get("r", ""))
        formulas.add(coordinate)
        formula_nodes.append((coordinate, formula))
        if _EXTERNAL_FORMULA.search(formula.text or ""):
            external.add(coordinate)
            if shared_id := formula.get("si"):
                shared_external.add(shared_id)
    for coordinate, formula in formula_nodes:
        if formula.get("si") in shared_external:
            external.add(coordinate)
    return frozenset(formulas), frozenset(external)


def _shape(package: OOXMLPackage, sheet: _Sheet, limit: int) -> _SheetShape:
    root = package.xml(sheet.part)
    if root is None:
        raise InvalidOOXMLError(f"XLSX worksheet XML is missing: {sheet.index}")
    cancel_check = package.cancel_check
    coordinates = [
        _coordinate(node.get("r", ""))
        for node in checked(root.iter(f"{{{_SHEET_NS}}}c"), cancel_check)
    ]
    hidden_rows = frozenset(
        int(node.get("r", "0"))
        for node in checked(root.iter(f"{{{_SHEET_NS}}}row"), cancel_check)
        if node.get("hidden") in {"1", "true", "True"} and node.get("r", "").isdecimal()
    )
    hidden_columns: set[int] = set()
    for node in checked(root.iter(f"{{{_SHEET_NS}}}col"), cancel_check):
        if node.get("hidden") not in {"1", "true", "True"}:
            continue
        try:
            minimum, maximum = int(node.get("min", "0")), int(node.get("max", "0"))
        except ValueError:
            raise InvalidOOXMLError(
                f"Invalid hidden-column range in XLSX sheet {sheet.index}"
            ) from None
        if minimum < 1 or maximum < minimum or maximum > 16_384:
            raise InvalidOOXMLError(f"Invalid hidden-column range in XLSX sheet {sheet.index}")
        hidden_columns.update(range(minimum, maximum + 1))
    merges = _merged_ranges(root, sheet, limit, cancel_check)
    formulas, external_formulas = _formula_cells(root, cancel_check)
    bounds = [*coordinates]
    bounds.extend((merge.min_row, merge.min_col) for merge in merges)
    bounds.extend((merge.max_row, merge.max_col) for merge in merges)
    if bounds:
        min_row = min(row for row, _ in bounds)
        max_row = max(row for row, _ in bounds)
        min_col = min(column for _, column in bounds)
        max_col = max(column for _, column in bounds)
    else:
        min_row = max_row = min_col = max_col = 0
    shape = _SheetShape(
        min_row,
        max_row,
        min_col,
        max_col,
        hidden_rows,
        frozenset(hidden_columns),
        bool(bounds),
        merges,
        formulas,
        external_formulas,
    )
    if shape.cells > limit:
        raise ResourceLimitError(
            f"XLSX sheet {sheet.index} exceeds the {limit}-cell used-range limit; split the sheet"
        )
    return shape


def _shapes(
    package: OOXMLPackage, selected: list[_Sheet], settings: Settings
) -> dict[str, _SheetShape]:
    """Measure every selected sheet, holding their combined area to one budget.

    Each used-range cell becomes a table cell whether or not it holds a value,
    so two corner cells per sheet would otherwise expand a small workbook into
    the per-sheet limit times the sheet count.
    """
    shapes: dict[str, _SheetShape] = {}
    total = 0
    for sheet in selected:
        shape = _shape(package, sheet, settings.limits.max_sheet_cells)
        total += shape.cells
        if total > settings.limits.max_workbook_cells:
            raise ResourceLimitError(
                f"XLSX selected sheets exceed the {settings.limits.max_workbook_cells}-cell "
                "combined used-range limit; select fewer sheets or split the workbook"
            )
        shapes[sheet.name] = shape
    return shapes


def _timespec(microsecond: int) -> str:
    """Keep the fraction a stored time carries; whole seconds print as before.

    openpyxl rounds Excel's serial times to milliseconds, so a fraction is
    normally three digits; a finer one is kept rather than rounded away.
    """
    if not microsecond:
        return "seconds"
    return "milliseconds" if microsecond % 1000 == 0 else "microseconds"


def _duration(value: dt.timedelta) -> str:
    # Whole microseconds, so a negative duration under one second keeps its sign.
    total = value // dt.timedelta(microseconds=1)
    sign = "-" if total < 0 else ""
    seconds_total, microseconds = divmod(abs(total), 1_000_000)
    hours, remainder = divmod(seconds_total, 3600)
    minutes, seconds = divmod(remainder, 60)
    fraction = ""
    if microseconds:
        digits = f"{microseconds:06d}"
        fraction = "." + (digits[:3] if _timespec(microseconds) == "milliseconds" else digits)
    return f"{sign}{hours}:{minutes:02d}:{seconds:02d}{fraction}"


def _temporal(value: object, number_format: str) -> str | None:
    if isinstance(value, dt.datetime):
        kind = is_datetime(number_format)
        if kind == "date":
            return value.date().isoformat()
        timespec = _timespec(value.microsecond)
        if kind == "time":
            return value.time().isoformat(timespec=timespec)
        return value.isoformat(sep=" ", timespec=timespec)
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat(timespec=_timespec(value.microsecond))
    if isinstance(value, dt.timedelta):
        return _duration(value)
    return None


def _display(
    cell: ReadOnlyCell | EmptyCell,
    counters: Counter[str],
    unsupported_formats: Counter[str],
    *,
    formula: bool = False,
    external_formula: bool = False,
) -> str:
    value = getattr(cell, "value", None)
    if external_formula:
        counters["external_references"] += 1
        return (
            '<span class="img-missing" role="note"><strong>'
            "External workbook value not loaded.</strong></span>"
        )
    if formula:
        counters["formulas"] += 1
        if value is None:
            return (
                '<span class="img-missing" role="note"><strong>'
                "Formula value not available.</strong></span>"
            )
    if value is None:
        return ""
    if (temporal := _temporal(value, getattr(cell, "number_format", "General"))) is not None:
        return html.escape(temporal)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    code = getattr(cell, "number_format", None) or "General"
    if isinstance(value, (int, float)):
        return html.escape(_formatted_number(value, code, unsupported_formats))
    text = _formatted_text(str(value), code, unsupported_formats)
    text = text.replace("\r\n", "\n").replace("\r", "\n").expandtabs(4)
    return html.escape(text).replace("\n", "<br>")


def _formatted_number(value: int | float, code: str, unsupported: Counter[str]) -> str:
    if code.casefold() == "general" or (isinstance(value, float) and not math.isfinite(value)):
        return general(value)
    rendered = format_number(value, code)
    if rendered is None:
        # Show the stored value unrounded rather than a guess at the code.
        unsupported[code] += 1
        return general(value)
    return rendered


def _formatted_text(text: str, code: str, unsupported: Counter[str]) -> str:
    if code.casefold() == "general":
        return text
    formatted = format_text(text, code)
    if formatted is None:
        unsupported[code] += 1
        return text
    return formatted


def _display_width(fragment: str) -> int:
    text = html.unescape(re.sub(r"<[^>]+>", "", fragment))
    return max((len(line) for line in text.splitlines()), default=0)


def _merge_layout(
    shape: _SheetShape, visible_rows: list[int], visible_columns: list[int]
) -> tuple[dict[tuple[int, int], tuple[int, int]], dict[tuple[int, int], tuple[int, int]]]:
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    covered: dict[tuple[int, int], tuple[int, int]] = {}
    header_row = visible_rows[0] if shape.header and visible_rows else None
    for merge in shape.merges:
        # The authored value resides in the top-left cell. Moving it out of a
        # hidden row or column would disclose content the author hid.
        if merge.min_row in shape.hidden_rows or merge.min_col in shape.hidden_columns:
            continue
        rows = [row for row in visible_rows if merge.min_row <= row <= merge.max_row]
        columns = [column for column in visible_columns if merge.min_col <= column <= merge.max_col]
        if not rows or not columns:
            continue
        # A cell cannot span from thead into tbody in valid table HTML. Leave
        # that unusual merge visually unmerged, retaining the anchor value.
        if header_row is not None and header_row in rows and len(rows) > 1:
            continue
        anchor = (merge.min_row, merge.min_col)
        if anchor != (rows[0], columns[0]):
            continue
        cells = [(row, column) for row in rows for column in columns]
        if any(cell in covered for cell in cells):
            raise InvalidOOXMLError("XLSX worksheet contains overlapping merged-cell ranges")
        anchors[anchor] = (len(rows), len(columns))
        covered.update(dict.fromkeys(cells, anchor))
    return anchors, covered


def _render_row(
    row_index: int,
    row: tuple[Any, ...],
    shape: _SheetShape,
    header_row: int | None,
    anchors: dict[tuple[int, int], tuple[int, int]],
    covered: dict[tuple[int, int], tuple[int, int]],
    widths: list[int],
    counters: Counter[str],
    unsupported_formats: Counter[str],
) -> str:
    cells: list[str] = []
    visible_index = 0
    for column, cell in enumerate(row, shape.min_col):
        if column in shape.hidden_columns:
            continue
        coordinate = (row_index, column)
        anchor = covered.get(coordinate)
        if anchor is not None and anchor != coordinate:
            visible_index += 1
            continue
        value = _display(
            cell,
            counters,
            unsupported_formats,
            formula=coordinate in shape.formulas,
            external_formula=coordinate in shape.external_formulas,
        )
        widths[visible_index] = max(widths[visible_index], _display_width(value))
        visible_index += 1
        tag = "th" if row_index == header_row else "td"
        attributes = ['scope="col"'] if tag == "th" else []
        if coordinate in anchors:
            rowspan, colspan = anchors[coordinate]
            if rowspan > 1:
                attributes.append(f'rowspan="{rowspan}"')
            if colspan > 1:
                attributes.append(f'colspan="{colspan}"')
        attrs = " " + " ".join(attributes) if attributes else ""
        cells.append(f"<{tag}{attrs}>{value}</{tag}>")
    return "<tr>" + "".join(cells) + "</tr>"


def _table(
    worksheet: Any,
    sheet: _Sheet,
    shape: _SheetShape,
    counters: Counter[str],
    unsupported_formats: Counter[str],
    cancel_check: CancelCheck = never_cancelled,
) -> str:
    if not shape.cells:
        return '<p class="pf-empty-sheet">No visible cell content.</p>'
    row_indexes = [
        row for row in range(shape.min_row, shape.max_row + 1) if row not in shape.hidden_rows
    ]
    column_indexes = [
        column
        for column in range(shape.min_col, shape.max_col + 1)
        if column not in shape.hidden_columns
    ]
    if not row_indexes or not column_indexes:
        return '<p class="pf-empty-sheet">No visible cell content.</p>'
    anchors, covered = _merge_layout(shape, row_indexes, column_indexes)
    widths = [0] * len(column_indexes)
    rendered_rows: list[str] = []
    header_row = row_indexes[0] if shape.header else None
    for row_index, row in zip(
        range(shape.min_row, shape.max_row + 1),
        worksheet.iter_rows(
            min_row=shape.min_row,
            max_row=shape.max_row,
            min_col=shape.min_col,
            max_col=shape.max_col,
        ),
        strict=True,
    ):
        cancel_check()
        if row_index in shape.hidden_rows:
            continue
        rendered_rows.append(
            _render_row(
                row_index,
                row,
                shape,
                header_row,
                anchors,
                covered,
                widths,
                counters,
                unsupported_formats,
            )
        )
    natural_width = sum(min(28, max(8, width)) * 7 + 16 for width in widths)
    thead = f"<thead>{rendered_rows[0]}</thead>" if header_row is not None else ""
    body_rows = rendered_rows[1:] if header_row is not None else rendered_rows
    tbody = "<tbody>" + "".join(body_rows) + "</tbody>"
    return (
        f'<table class="pf-data-table" data-natural-width="{natural_width}">'
        f"<caption>{html.escape(sheet.name)}</caption>{thead}{tbody}</table>"
    )


def _anchor_position(anchor: Any) -> tuple[int, int]:
    origin = anchor.find(f"{{{_DRAWING_NS}}}from")
    if origin is None:
        return 1, 1
    try:
        row = int(origin.findtext(f"{{{_DRAWING_NS}}}row") or "0") + 1
        column = int(origin.findtext(f"{{{_DRAWING_NS}}}col") or "0") + 1
    except ValueError:
        raise InvalidOOXMLError("Invalid XLSX drawing anchor") from None
    return row, column


def _embed_image(
    package: OOXMLPackage, image_part: str, suffix: str, assets: dict[str, bytes]
) -> str:
    """Carry one image part once, however many anchors or sheets place it.

    Reading the part per anchor would let one large image referenced from
    thousands of anchors multiply into gigabytes before any later limit runs.
    """
    digest = hashlib.sha256(image_part.encode("utf-8")).hexdigest()[:16]
    key = f"__platefold_xlsx_asset_{digest}{suffix}"
    if key in assets:
        return key
    raw = package.read(image_part, max_bytes=package.limits.max_image_bytes)
    limit = package.limits.max_embedded_bytes
    if sum(len(value) for value in assets.values()) + len(raw) > limit:
        raise ResourceLimitError(f"XLSX images exceed the {limit}-byte embedded asset limit")
    assets[key] = raw
    return key


def _drawing_fragment(
    package: OOXMLPackage,
    anchor: Any,
    relationships: dict[str, tuple[str, str, bool]],
    shape: _SheetShape,
    assets: dict[str, bytes],
    counters: Counter[str],
) -> str:
    row, column = _anchor_position(anchor)
    if row in shape.hidden_rows or column in shape.hidden_columns:
        counters["hidden_drawings"] += 1
        return ""
    blip = anchor.find(f".//{{{_DRAWINGML_NS}}}blip")
    chart = anchor.find(f".//{{{_CHART_NS}}}chart")
    if blip is not None:
        linked = blip.get(f"{{{_DOC_REL_NS}}}link")
        embedded = blip.get(f"{{{_DOC_REL_NS}}}embed")
        relation = relationships.get(linked or embedded or "")
        if linked or relation is None or relation[2]:
            counters["linked_images"] += 1
            return (
                '<div class="img-missing" role="note"><strong>'
                "Linked spreadsheet image not loaded.</strong></div>"
            )
        relation_kind, image_part, _ = relation
        suffix = Path(image_part).suffix.lower()
        if relation_kind != "image" or suffix not in _IMAGE_SUFFIXES:
            counters["unsupported_images"] += 1
            return (
                '<div class="img-missing" role="note"><strong>'
                "Spreadsheet image format is not supported.</strong></div>"
            )
        key = _embed_image(package, image_part, suffix, assets)
        properties = anchor.find(f".//{{{_DRAWING_NS}}}cNvPr")
        alternative = (properties.get("descr", "") if properties is not None else "").strip()
        if alternative.casefold() in {"image", "picture"}:
            alternative = ""
        if not alternative:
            counters["images_without_alt"] += 1
            alternative = "Spreadsheet image without authored alternative text."
        return f'<img src="{key}" alt="{html.escape(alternative, quote=True)}">'
    if chart is not None:
        counters["charts"] += 1
        return (
            '<div class="img-missing" role="note"><strong>'
            "Excel chart not rendered.</strong> Add a static image to include its visual content."
            "</div>"
        )
    counters["unsupported_drawings"] += 1
    return (
        '<div class="img-missing" role="note"><strong>'
        "Spreadsheet drawing not rendered.</strong></div>"
    )


def _drawings(
    package: OOXMLPackage,
    sheet: _Sheet,
    shape: _SheetShape,
    assets: dict[str, bytes],
    counters: Counter[str],
) -> str:
    sheet_relationships = _relationships(package, sheet.part)
    fragments: list[str] = []
    for kind, drawing_part, external in sheet_relationships.values():
        if kind != "drawing" or external:
            continue
        drawing = package.xml(drawing_part)
        if drawing is None:
            raise InvalidOOXMLError(f"XLSX drawing part is missing for sheet {sheet.index}")
        drawing_relationships = _relationships(package, drawing_part)
        for anchor in checked(list(drawing), package.cancel_check, 1):
            fragments.append(
                _drawing_fragment(package, anchor, drawing_relationships, shape, assets, counters)
            )
    return '<div class="pf-sheet-drawings">' + "".join(fragments) + "</div>" if fragments else ""


def _embedded_objects(
    package: OOXMLPackage, sheet: _Sheet, shape: _SheetShape, counters: Counter[str]
) -> str:
    root = package.xml(sheet.part)
    assert root is not None
    visible = 0
    for node in checked(root.iter(f"{{{_SHEET_NS}}}oleObject"), package.cancel_check):
        position = _generic_anchor_position(node)
        if position is not None and (
            position[0] in shape.hidden_rows or position[1] in shape.hidden_columns
        ):
            counters["hidden_drawings"] += 1
            continue
        visible += 1
    counters["external_objects"] += visible
    if not visible:
        return ""
    return (
        '<div class="img-missing" role="note"><strong>'
        "Embedded spreadsheet object not loaded.</strong></div>"
    ) * visible


def _generic_anchor_position(node: Any) -> tuple[int, int] | None:
    origin = next((child for child in node.iter() if child.tag.rsplit("}", 1)[-1] == "from"), None)
    if origin is None:
        return None
    values = {child.tag.rsplit("}", 1)[-1]: child.text for child in origin}
    try:
        return int(values.get("row") or "0") + 1, int(values.get("col") or "0") + 1
    except ValueError:
        raise InvalidOOXMLError("Invalid XLSX object anchor") from None


def _hyperlink_count(package: OOXMLPackage, sheet: _Sheet, shape: _SheetShape) -> int:
    root = package.xml(sheet.part)
    assert root is not None
    count = 0
    for node in checked(root.iter(f"{{{_SHEET_NS}}}hyperlink"), package.cancel_check):
        try:
            min_col, min_row, _max_col, _max_row = range_boundaries(node.get("ref", ""))
        except (TypeError, ValueError):
            raise InvalidOOXMLError("Invalid XLSX hyperlink range") from None
        if not isinstance(min_row, int) or not isinstance(min_col, int):
            raise InvalidOOXMLError("Invalid XLSX hyperlink range")
        if min_row not in shape.hidden_rows and min_col not in shape.hidden_columns:
            count += 1
    return count


def _metadata(workbook: Any) -> tuple[str | None, list[tuple[str, str]]]:
    properties = workbook.properties
    title = (properties.title or "").strip() or None
    metadata = []
    for label, value in (
        ("Author", properties.creator),
        ("Version", properties.version),
        ("Document", properties.subject),
    ):
        text = str(value or "").strip()
        if text:
            metadata.append((label, text))
    return title, metadata


def _format_warning(unsupported_formats: Counter[str]) -> tuple[str, str]:
    """Name the most-used unsupported codes, each bounded in length."""
    shown = [
        code if len(code) <= _FORMAT_NAME_LIMIT else code[: _FORMAT_NAME_LIMIT - 3] + "..."
        for code, _ in unsupported_formats.most_common(_FORMATS_NAMED)
    ]
    names = ", ".join(json.dumps(code, ensure_ascii=False) for code in shown)
    more = len(unsupported_formats) - len(shown)
    if more:
        names += f" and {more} more"
    cells = sum(unsupported_formats.values())
    return (
        "XLSX_NUMBER_FORMAT_UNSUPPORTED",
        f"{cells} cell(s) use Excel number formats that Plate|Fold does not render and show "
        f"their stored values unformatted instead: {names}.",
    )


def _warnings(
    counters: Counter[str], unsupported_formats: Counter[str] | None = None
) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    hidden = {
        key: counters[key]
        for key in ("hidden_sheets", "hidden_rows", "hidden_columns", "hidden_drawings")
    }
    if any(hidden.values()):
        result.append(
            (
                "XLSX_HIDDEN_CONTENT",
                "Skipped hidden content: "
                f"{hidden['hidden_sheets']} sheet(s), {hidden['hidden_rows']} row(s), "
                f"{hidden['hidden_columns']} column(s), {hidden['hidden_drawings']} drawing(s).",
            )
        )
    if unsupported_formats:
        result.append(_format_warning(unsupported_formats))
    messages = (
        (
            "external_references",
            "XLSX_EXTERNAL_REFERENCE",
            "external cell reference(s) were not loaded",
        ),
        ("linked_images", "XLSX_LINKED_IMAGE", "linked image(s) were not loaded"),
        ("hyperlinks", "XLSX_HYPERLINK_SKIPPED", "hyperlink(s) were rendered as plain text"),
        (
            "formulas",
            "XLSX_FORMULA_UNCALCULATED",
            "formula cell(s) used stored values when available and were not recalculated",
        ),
        ("charts", "XLSX_CHART_UNSUPPORTED", "Excel chart(s) were replaced with visible warnings"),
        (
            "unsupported_images",
            "XLSX_IMAGE_UNSUPPORTED",
            "unsupported embedded image(s) were replaced with visible warnings",
        ),
        (
            "images_without_alt",
            "XLSX_IMAGE_ALT_MISSING",
            "embedded image(s) had no authored alternative text",
        ),
        (
            "unsupported_drawings",
            "XLSX_DRAWING_UNSUPPORTED",
            "unsupported drawing object(s) were replaced with visible warnings",
        ),
        (
            "chart_sheets",
            "XLSX_CHART_SHEET_UNSUPPORTED",
            "visible chart sheet(s) were not exported",
        ),
        (
            "external_objects",
            "XLSX_OBJECT_SKIPPED",
            "OLE object(s) were not loaded",
        ),
        (
            "library_warnings",
            "XLSX_LIBRARY_WARNING",
            "workbook feature warning(s) were reported by the XLSX reader",
        ),
    )
    for key, code, detail in messages:
        if counters[key]:
            result.append((code, f"{counters[key]} {detail}."))
    return result


def source_theme_for_path(path: str | Path, settings: Settings):
    """Read only the reusable DrawingML theme from one validated XLSX."""
    source = Path(path)
    with OOXMLPackage(
        source,
        settings.limits,
        max_xml_nodes=settings.limits.max_xlsx_xml_nodes,
        allowed_external_relationships=_ALLOWED_EXTERNAL_RELATIONSHIPS,
    ) as package:
        if not package.has("xl/workbook.xml"):
            raise InvalidOOXMLError("XLSX package has no xl/workbook.xml part")
        return drawingml_theme(package, "xl/theme/theme1.xml", label=source.name)


def ingest(path: str, settings: Settings, cancel_check: CancelCheck | None = None) -> Ingested:
    source = Path(path)
    cancel_check = cancel_check or never_cancelled
    assets: dict[str, bytes] = {}
    counters: Counter[str] = Counter()
    unsupported_formats: Counter[str] = Counter()
    try:
        with OOXMLPackage(
            source,
            settings.limits,
            max_xml_nodes=settings.limits.max_xlsx_xml_nodes,
            allowed_external_relationships=_ALLOWED_EXTERNAL_RELATIONSHIPS,
            cancel_check=cancel_check,
        ) as package:
            sheets = _sheets(package)
            if len(sheets) > settings.limits.max_sheet_count:
                raise ResourceLimitError(
                    f"XLSX workbook exceeds the {settings.limits.max_sheet_count}-sheet limit"
                )
            selected = _selected_sheets(sheets, settings.xlsx_sheets)
            if not selected:
                raise ValueError("XLSX workbook contains no visible worksheets")
            counters["hidden_sheets"] = sum(sheet.state != "visible" for sheet in sheets)
            counters["chart_sheets"] = sum(
                sheet.state == "visible" and not sheet.part.startswith("xl/worksheets/")
                for sheet in sheets
            )
            theme, theme_warnings = drawingml_theme(
                package, "xl/theme/theme1.xml", label=source.name
            )
            shapes = _shapes(package, selected, settings)
            cancel_check()
            with python_warnings.catch_warnings(record=True) as notices:
                python_warnings.simplefilter("always")
                workbook = load_workbook(
                    source,
                    read_only=True,
                    keep_vba=False,
                    data_only=True,
                    keep_links=False,
                    rich_text=False,
                )
            counters["library_warnings"] += len(notices)
            try:
                title, metadata = _metadata(workbook)
                language = (workbook.properties.language or "").strip() or None
                ids: set[str] = set()
                outline: list[tuple[int, str, str]] = []
                sections: list[str] = []
                for sheet in selected:
                    cancel_check()
                    shape = shapes[sheet.name]
                    counters["hidden_rows"] += sum(
                        shape.min_row <= row <= shape.max_row for row in shape.hidden_rows
                    )
                    counters["hidden_columns"] += sum(
                        shape.min_col <= column <= shape.max_col for column in shape.hidden_columns
                    )
                    anchor = unique(slugify(sheet.name, "-"), ids)
                    outline.append((1, sheet.name, anchor))
                    table = _table(
                        workbook[sheet.name],
                        sheet,
                        shape,
                        counters,
                        unsupported_formats,
                        cancel_check,
                    )
                    drawings = _drawings(package, sheet, shape, assets, counters)
                    objects = _embedded_objects(package, sheet, shape, counters)
                    counters["hyperlinks"] += _hyperlink_count(package, sheet, shape)
                    sections.append(
                        '<section class="pf-sheet-section">'
                        '<section class="pf-data-section">'
                        f'<h1 id="{html.escape(anchor, quote=True)}">{html.escape(sheet.name)}</h1>'
                        f"{table}</section>{drawings}{objects}</section>"
                    )
            finally:
                workbook.close()
    except (ResourceLimitError, InvalidOOXMLError, ValueError, CancelledError):
        raise
    except Exception as exc:
        raise InvalidOOXMLError(f"Invalid XLSX structure: {source.name}") from exc
    return Ingested(
        body_html="".join(sections),
        title=title or source.stem,
        language=language,
        metadata=metadata,
        outline=outline,
        assets=assets,
        source_theme=theme,
        warnings=[*theme_warnings, *_warnings(counters, unsupported_formats)],
    )
