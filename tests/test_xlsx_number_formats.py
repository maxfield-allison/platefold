"""Excel number formats and sheet selectors through the real XLSX ingest path."""

from __future__ import annotations

import html
import json
import re

import pytest
from openpyxl import Workbook
from pypdf import PdfReader

from platefold.cli import build_settings, parse_args
from platefold.config import Settings, make_settings, parse_settings_file
from platefold.ingest import ingest
from platefold.server_models import JobLimits, JobRequest
from platefold.worker import request_from

# (label, stored value, format code, text Excel displays)
FORMATTED = [
    # Fractions: the headline defect showed 0.25 under `# ?/?` as 0.
    ("fraction", 0.25, "# ?/?", "1/4"),
    ("mixed fraction", 1.25, "# ?/?", "1 1/4"),
    ("two-digit fraction", 0.3125, "# ??/??", "5/16"),
    ("three-digit fraction", 0.3333, "# ???/???", "1/3"),
    ("improper fraction", 1.25, "?/?", "5/4"),
    ("eighths", 0.5, "# ?/8", "4/8"),
    ("hundredths", 1.3, "# ?/100", "1 30/100"),
    ("whole under fraction", 2, "# ?/?", "2"),
    ("zero under fraction", 0, "# ?/?", "0"),
    ("zero whole placeholder", 0.25, "0 ?/?", "0 1/4"),
    ("negative fraction", -0.25, "# ?/?", "-1/4"),
    # Zero padding keeps identifiers intact.
    ("identifier", 123, "000000", "000123"),
    ("social security layout", 123456789, "000-00-0000", "123-45-6789"),
    ("phone layout", 5551234, "(###) ###-####", "() 555-1234"),
    # Separators, decimals, percent and rounding half away from zero.
    ("thousands", 1234567.891, "#,##0.00", "1,234,567.89"),
    ("half away from zero", 0.125, "0.00", "0.13"),
    ("binary half", 2.675, "0.00", "2.68"),
    ("optional decimals", 1.5, "0.0#", "1.5"),
    ("percent", 0.125, "0.00%", "12.50%"),
    ("no integer placeholder", 0.5, "#.00", ".50"),
    # Scaling commas and their units.
    ("scaled millions", 1250000, '0.0,," million"', "1.3 million"),
    ("scaled thousands", 1234567, '#,##0,"K"', "1,235K"),
    # Literals are text, not directives.
    ("literal percent", 0.25, '0.00"%"', "0.25%"),
    ("escaped plus", 3, "\\+0", "+3"),
    ("only literal", 5, '"Total"', "Total"),
    # Sections: positive; negative; zero; text.
    ("negative section", -1234.5, "#,##0.00;(#,##0.00)", "(1,234.50)"),
    ("single-section negative", -1234.5, "$#,##0.00", "-$1,234.50"),
    ("zero section", 0, '0.00;(0.00);"zero"', "zero"),
    ("empty negative section", -5, "0;;", ""),
    (
        "accounting positive",
        1234.5,
        '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)',
        "$ 1,234.50",
    ),
    ("accounting negative", -12, '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', "$ (12.00)"),
    ("accounting zero", 0, '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)', "$ -"),
    # Conditions choose the section.
    ("condition second", 50, "[>=100]0.0;[<100]0.00", "50.00"),
    ("condition first", 150, "[>=100]0.0;[<100]0.00", "150.0"),
    ("condition fallback", 7, '[=1]"one";"other "0', "other 7"),
    # Currency, locale and colour brackets.
    ("euro locale", 1234.5, "[$€-407] #,##0.00", "€ 1,234.50"),
    ("locale only", 42, "[$-409]0", "42"),
    ("colours", -2, "[Blue]0.00;[Red](0.00)", "(2.00)"),
    # Scientific notation and General inside a section.
    ("scientific", 12345, "0.00E+00", "1.23E+04"),
    ("small scientific", 0.00012345, "0.00E+00", "1.23E-04"),
    ("general with literal", 5, 'General" units"', "5 units"),
]

TEXT_FORMATTED = [
    ("text placeholder", "abc", '"ID-"@', "ID-abc"),
    ("four-section text", "abc", '0;-0;0;"<"@">"', "<abc>"),
    ("numeric-only code on text", "abc", "0.00", "abc"),
    # A text cell never uses the numeric sections, so a date code is no reason to warn.
    ("text under a date code", "N/A", "yyyy-mm-dd", "N/A"),
    ("hidden text section", "secret", "0;-0;0;", ""),
]

# (label, stored value, format code, unrounded text shown instead)
UNSUPPORTED = [
    ("engineering", 12345.678, "##0.0E+0", "12345.678"),
    ("negative under a condition", -5.25, "[<0]0;0", "-5.25"),
    ("double-byte digits", 1234.5, "[DBNum1]0", "1234.5"),
    ("scaling before the point", 1234567, "0,.00", "1234567"),
    ("two percent signs", 0.5, "0%%", "0.5"),
]


def _rows(body: str) -> dict[str, str]:
    """Map each data row's label to its rendered value cell."""
    cells = re.findall(r"<tr><td>(.*?)</td><td>(.*?)</td></tr>", body)
    return {html.unescape(label): html.unescape(value) for label, value in cells}


def _format_book(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formats"
    sheet.append(["Label", "Value"])
    for label, value, code, _expected in rows:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 2).number_format = code
    workbook.save(path)
    return path


def test_supported_number_formats_show_what_excel_shows(tmp_path):
    path = _format_book(tmp_path / "formats.xlsx", FORMATTED + TEXT_FORMATTED)
    result = ingest(str(path), Settings())
    rendered = _rows(result.body_html)
    expected = {label: text for label, _value, _code, text in FORMATTED + TEXT_FORMATTED}
    assert rendered == expected
    assert not [code for code, _ in result.warnings if code == "XLSX_NUMBER_FORMAT_UNSUPPORTED"]


def test_unsupported_formats_keep_the_stored_value_and_report_each_code(tmp_path):
    rows = [*UNSUPPORTED, ("text section digits", "kept", "0;0;0;0", "kept")]
    path = _format_book(tmp_path / "unsupported.xlsx", rows)
    result = ingest(str(path), Settings())
    rendered = _rows(result.body_html)
    assert rendered == {label: text for label, _value, _code, text in rows}
    details = [
        detail for code, detail in result.warnings if code == "XLSX_NUMBER_FORMAT_UNSUPPORTED"
    ]
    assert len(details) == 1
    assert details[0].startswith("6 cell(s) use Excel number formats")
    # Three codes are named and the rest are counted, each quoted as authored.
    assert details[0].count('"') == 6 and details[0].endswith(" and 3 more.")


def test_dates_keep_their_existing_iso_rendering(tmp_path):
    import datetime as dt

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Label", "Value"])
    sheet.append(["date", dt.date(2026, 9, 14)])
    sheet["B2"].number_format = "d mmmm yyyy"
    sheet.append(["time", dt.time(14, 30)])
    sheet["B3"].number_format = "h:mm AM/PM"
    path = tmp_path / "dates.xlsx"
    workbook.save(path)
    result = ingest(str(path), Settings())
    assert _rows(result.body_html) == {"date": "2026-09-14", "time": "14:30:00"}
    assert not result.warnings or all(
        code != "XLSX_NUMBER_FORMAT_UNSUPPORTED" for code, _ in result.warnings
    )


def test_fractional_seconds_survive_in_times_datetimes_and_durations(tmp_path):
    import datetime as dt

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Label", "Value"])
    cells = [
        ("time", dt.time(12, 34, 56, 789_000), "hh:mm:ss.000"),
        ("datetime", dt.datetime(2026, 9, 14, 8, 15, 30, 250_000), "yyyy-mm-dd hh:mm:ss.000"),
        ("duration", dt.timedelta(hours=1, minutes=2, seconds=3, milliseconds=500), "[h]:mm:ss.0"),
        ("negative subsecond", dt.timedelta(milliseconds=-250), "[h]:mm:ss.00"),
        ("whole seconds", dt.time(14, 30, 5), "hh:mm:ss"),
    ]
    for label, value, code in cells:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 2).number_format = code
    path = tmp_path / "fractions.xlsx"
    workbook.save(path)

    result = ingest(str(path), Settings())

    assert _rows(result.body_html) == {
        "time": "12:34:56.789",
        "datetime": "2026-09-14 08:15:30.250",
        "duration": "1:02:03.500",
        "negative subsecond": "-0:00:00.250",
        "whole seconds": "14:30:05",
    }


@pytest.mark.integration
def test_real_export_prints_formatted_values_and_reports_unsupported_codes(tmp_path, run_export):
    rows = [
        ("fraction", 0.25, "# ?/?", "1/4"),
        ("identifier", 123, "000000", "000123"),
        ("scaled", 1250000, '0.0,," million"', "1.3 million"),
        ("literal", 0.25, '0.00"%"', "0.25%"),
        ("conditional", 50, "[>=100]0.0;[<100]0.00", "50.00"),
        ("negative", -1234.5, "#,##0.00;(#,##0.00)", "(1,234.50)"),
        ("zero", 0, '0.00;(0.00);"nil"', "nil"),
        ("engineering", 12345.678, "##0.0E+0", "12345.678"),
    ]
    path = _format_book(tmp_path / "printed.xlsx", rows)
    code, output = run_export(path, "--no-toc")
    assert code == 0
    reader = PdfReader(output / "printed.pdf")
    text = "".join("".join(page.extract_text() or "" for page in reader.pages).split())
    for _label, _value, _code, expected in rows:
        assert "".join(expected.split()) in text
    assert "1250000" not in text
    report = json.loads((output / "_export-report.json").read_text("utf-8"))
    warnings = report["documents"][0]["warnings"]
    assert any(
        warning.startswith("XLSX_NUMBER_FORMAT_UNSUPPORTED") and '"##0.0E+0"' in warning
        for warning in warnings
    )


def _selection_book(path):
    workbook = Workbook()
    names = ["#Summary", "Data", "#2", "2", "#9"]
    workbook.active.title = names[0]
    for name in names[1:]:
        workbook.create_sheet(name)
    for sheet in workbook.worksheets:
        sheet.append(["Sheet", "Marker"])
        sheet.append([sheet.title, f"marker-{sheet.title}"])
    workbook.save(path)
    return path


def _selected(path, selectors, settings=None):
    settings = settings or Settings(xlsx_sheets=tuple(selectors))
    return [name for _, name, _ in ingest(str(path), settings).outline]


@pytest.mark.parametrize(
    ("selector", "sheet"),
    [
        ("#Summary", "#Summary"),  # a hash that is not a number is a name
        ("#2", "Data"),  # hash and a whole number is always a position
        ("name:#2", "#2"),  # the escape selects the literal name
        ("2", "2"),  # a bare number prefers a sheet of that name
        ("#4", "2"),
        ("name:2", "2"),
        ("Data", "Data"),
        ("1", "#Summary"),
    ],
)
def test_sheet_selectors_resolve_names_positions_and_the_name_escape(tmp_path, selector, sheet):
    path = _selection_book(tmp_path / "selectors.xlsx")
    assert _selected(path, [selector]) == [sheet]


@pytest.mark.parametrize("selector", ["#0", "#02", "#-1", "name:Missing", "name:", "#Missing"])
def test_selectors_that_are_neither_a_position_nor_a_name_are_refused(tmp_path, selector):
    path = _selection_book(tmp_path / "refused.xlsx")
    with pytest.raises(ValueError, match="Unknown XLSX sheet selection"):
        _selected(path, [selector])


def test_out_of_range_position_names_the_escape_for_a_matching_sheet(tmp_path):
    path = _selection_book(tmp_path / "hint.xlsx")
    with pytest.raises(ValueError, match="with name:#9"):
        _selected(path, ["#9"])
    assert _selected(path, ["name:#9"]) == ["#9"]


def test_cli_settings_file_http_and_desktop_select_the_same_sheets(tmp_path):
    path = _selection_book(tmp_path / "agree.xlsx")
    selectors = ["#Summary", "name:#2", "#2"]
    expected = ["#Summary", "#2", "Data"]

    args = parse_args([str(path), "--sheets", *selectors])
    cli_settings, _ = build_settings(args, str(tmp_path))
    _, overrides = parse_settings_file(
        "behaviour:\n  sheets: " + json.dumps(selectors, ensure_ascii=False) + "\n"
    )
    file_settings = make_settings(**overrides)
    http_settings = JobRequest.model_validate(
        {
            "source": {"kind": "upload", "format": "xlsx"},
            "settings": {"xlsx_sheets": selectors},
        }
    ).settings.settings(JobLimits())
    desktop_settings = request_from(
        {
            "input_path": str(path),
            "output_dir": str(tmp_path / "out"),
            "settings": {"xlsx_sheets": selectors},
        }
    ).settings

    for settings in (cli_settings, file_settings, http_settings, desktop_settings):
        assert settings.xlsx_sheets == tuple(selectors)
        assert _selected(path, selectors, settings) == expected
