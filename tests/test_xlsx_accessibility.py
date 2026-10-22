"""Combined XLSX/accessibility evidence; not native screen-reader acceptance."""

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from test_xlsx_ingest import _png, _workbook

from platefold.cli import build_settings, parse_args
from platefold.config import Settings, make_settings
from platefold.document import Document
from platefold.ingest import ingest
from platefold.service import ExportRequest, export
from tools.a11y.compare import raster_hashes
from tools.a11y.inspect_pdf import inspect


def workbook_source(path, language=" fr-CA ", rows=0):
    _workbook(path, wide=True)
    book = load_workbook(path)
    book.properties.language = language
    for row in range(rows):
        book["Summary ü"].append([f"R{row}C{column}" for column in range(60)])
    book["Details"].merge_cells("A3:B3")
    book["Details"]["A3"] = "Merged explanation"
    book["Details"].add_image(Image(_png()), "D5")
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    drawing = ElementTree.fromstring(parts["xl/drawings/drawing1.xml"])
    drawing.find(".//{*}cNvPr").set("descr", "A blue reference rectangle.")
    parts["xl/drawings/drawing1.xml"] = ElementTree.tostring(drawing)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


@pytest.mark.parametrize(
    "language,expected", [(" fr-CA ", "fr-CA"), (None, "und"), ("en_US", "und")]
)
def test_xlsx_language_validation_and_explicit_override(tmp_path, language, expected):
    path = workbook_source(tmp_path / "source.xlsx", language)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    assert ingest(str(path), Settings()).language == ((language or "").strip() or None)
    body, info = Document(
        str(path), str(tmp_path), make_settings(locale_hint="ja-JP"), None
    ).render()
    assert f'<html lang="{expected}">' in body and info["language"] == expected
    settings, _ = build_settings(
        parse_args([str(path), "--sheets", "Details", "--document-language", "de-DE"]),
        str(tmp_path),
    )
    body, info = Document(str(path), str(tmp_path), settings, None).render()
    assert '<html lang="de-DE">' in body and info["language"] == "de-DE"
    assert "Summary ü" not in body and "A blue reference rectangle." in body
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


@pytest.mark.integration
@pytest.mark.parametrize(
    "mode,override,rows",
    [("online", "", 0), ("print", "de-DE", 0), ("tiled", "", 0), ("print", "", 100)],
)
def test_xlsx_final_pdf_preserves_semantics_and_bookmark_rewrite(
    tmp_path, monkeypatch, mode, override, rows
):
    from platefold import service

    source = workbook_source(tmp_path / "source.xlsx", rows=rows)
    original = service.print_pdf
    captures = []

    def capture(page, html, destination, *args, **options):
        original(page, html, destination, *args, **options)
        raw = tmp_path / f"raw-{len(captures)}.pdf"
        raw.write_bytes(Path(destination).read_bytes())
        captures.append(raw)

    monkeypatch.setattr(service, "print_pdf", capture)
    result = export(
        ExportRequest(
            str(source),
            str(tmp_path / "final"),
            settings=make_settings(
                output_mode=mode, document_language=override, locale_hint="ja-JP"
            ),
            browser="bundled",
        )
    )
    assert result.status == "completed_with_warnings"  # Authored hidden rows/sheets are omitted.
    final = tmp_path / "final" / "source.pdf"
    before, after = inspect(captures[-1]), inspect(final)
    assert after["language"] == (override or "fr-CA")
    assert after["structure_tree"] and after["parent_tree"] and not after["parent_tree_errors"]
    assert after["nodes"] == before["nodes"]
    assert raster_hashes(final) == raster_hashes(captures[-1])
    assert after["roles"]["/TH"] >= 61  # 59 visible wide-sheet columns plus two details columns.
    assert any("/Headers" in str(node.get("/A")) for node in after["nodes"])
    assert any("/ColSpan" in str(node.get("/A")) for node in after["nodes"])
    assert after["alternatives"] == ["A blue reference rectangle."]
    assert after["outline_top_level"] >= 2
    text = "\n".join(after["text"])
    assert text.index("Summary ü") < text.index("Details")
    assert "HIDDEN-" not in text and "Private" not in text

    if rows:
        assert after["pages"] > 3
        assert f"R{rows - 1}C59" in "".join(text.split())
        assert after["roles"]["/TD"] >= 59 * (rows + 1) + 3
