"""XLSX sheet policy, OOXML safety, drawings and real wide-page output."""

from __future__ import annotations

import datetime as dt
import json
import zipfile
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.table import Table
from PIL import Image
from pypdf import PdfReader

from platefold.cli import build_settings, parse_args
from platefold.config import Settings, make_settings, parse_settings_file
from platefold.document import Document
from platefold.ingest import SUPPORTED_SUFFIXES, ingest
from platefold.limits import ResourceLimitError, ResourceLimits
from platefold.ooxml import InvalidOOXMLError
from platefold.service import ExportRequest, _save_requested_theme, collect, prepare_export


class NoDiagrams:
    def render(self, _source):
        raise AssertionError("XLSX ingestion must not invoke Mermaid")


def _png(colour: str = "#214e75") -> BytesIO:
    output = BytesIO()
    Image.new("RGB", (24, 12), colour).save(output, "PNG")
    output.seek(0)
    return output


def _workbook(path, *, wide: bool = False):
    workbook = Workbook()
    workbook.properties.title = "Quarterly workbook"
    workbook.properties.creator = "PlateFold Test"
    summary = workbook.active
    summary.title = "Summary ü"
    summary.freeze_panes = "A2"
    columns = 60 if wide else 5
    summary.append([f"Column {index}" for index in range(1, columns + 1)])
    row = ["Alpha", 0.125, dt.date(2026, 9, 14), 1234.5, "مرحبا"]
    summary.append((row + list(range(columns)))[:columns])
    summary["B2"].number_format = "0.00%"
    summary["C2"].number_format = "yyyy-mm-dd"
    summary["D2"].number_format = "$#,##0.00"
    summary["A3"] = "HIDDEN-ROW-SENTINEL"
    summary.row_dimensions[3].hidden = True
    summary["E2"] = "HIDDEN-COLUMN-SENTINEL"
    summary.column_dimensions["E"].hidden = True

    details = workbook.create_sheet("Details")
    details.append(["Name", "Value"])
    details.append(["Beta", 7])
    details.add_table(Table(displayName="DetailsTable", ref="A1:B2"))

    hidden = workbook.create_sheet("Private")
    hidden["A1"] = "HIDDEN-SHEET-SENTINEL"
    hidden.sheet_state = "hidden"
    workbook.save(path)
    return path


def _warning_codes(result) -> set[str]:
    return {code for code, _ in result.warnings}


def test_registry_cli_settings_and_collection_accept_xlsx(tmp_path):
    path = _workbook(tmp_path / "book.XLSX")
    assert ".xlsx" in SUPPORTED_SUFFIXES
    files, _ = collect(str(tmp_path))
    assert files == [str(path)]

    args = parse_args([str(path), "--sheets", "Details", "1"])
    settings, config_path = build_settings(args, str(tmp_path))
    assert settings.xlsx_sheets == ("Details", "1") and config_path is None

    _, overrides = parse_settings_file('behaviour:\n  sheets: ["Summary ü", "#2"]\n')
    assert make_settings(**overrides).xlsx_sheets == ("Summary ü", "#2")


def test_visible_sections_formats_headers_and_hidden_content(tmp_path):
    path = _workbook(tmp_path / "proof.xlsx")
    result = ingest(str(path), Settings())

    assert result.title == "Quarterly workbook"
    assert ("Author", "PlateFold Test") in result.metadata
    assert result.outline == [(1, "Summary ü", "summary-u"), (1, "Details", "details")]
    assert result.body_html.count('class="pf-data-table"') == 2
    assert '<th scope="col">Column 1</th>' in result.body_html
    assert '<th scope="col">Name</th>' in result.body_html
    assert "12.50%" in result.body_html
    assert "2026-09-14" in result.body_html and "00:00:00" not in result.body_html
    assert "$1,234.50" in result.body_html
    assert "مرحبا" not in result.body_html  # its entire column is hidden
    assert "HIDDEN-ROW-SENTINEL" not in result.body_html
    assert "HIDDEN-COLUMN-SENTINEL" not in result.body_html
    assert "HIDDEN-SHEET-SENTINEL" not in result.body_html
    assert result.source_theme is not None
    hidden = next(detail for code, detail in result.warnings if code == "XLSX_HIDDEN_CONTENT")
    assert "1 sheet(s), 1 row(s), 1 column(s)" in hidden


def test_first_visible_used_row_is_the_header_without_excel_table_metadata(tmp_path):
    path = tmp_path / "plain.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Value"])
    sheet.append(["Alpha", 3])
    workbook.save(path)

    result = ingest(str(path), Settings())
    assert '<thead><tr><th scope="col">Name</th><th scope="col">Value</th></tr></thead>' in (
        result.body_html
    )
    assert "<tbody><tr><td>Alpha</td><td>3</td></tr></tbody>" in result.body_html


def test_sheet_selection_names_indexes_order_deduplication_and_hidden_refusal(tmp_path):
    path = _workbook(tmp_path / "selected.xlsx")
    by_name = ingest(str(path), Settings(xlsx_sheets=("Details", "Summary ü", "Details")))
    assert [name for _, name, _ in by_name.outline] == ["Details", "Summary ü"]
    by_index = ingest(str(path), Settings(xlsx_sheets=("#2", "1")))
    assert [name for _, name, _ in by_index.outline] == ["Details", "Summary ü"]

    with pytest.raises(ValueError, match="Unknown XLSX sheet selection"):
        ingest(str(path), Settings(xlsx_sheets=("Missing",)))
    with pytest.raises(ValueError, match="hidden and cannot be exported"):
        ingest(str(path), Settings(xlsx_sheets=("Private",)))
    with pytest.raises(ValueError, match="hidden and cannot be exported"):
        ingest(str(path), Settings(xlsx_sheets=("3",)))


def test_formulas_and_hyperlinks_are_not_evaluated_or_followed(tmp_path):
    path = tmp_path / "links.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Links"
    sheet.append(["Kind", "Value"])
    sheet.append(["Internal formula", "=SUM(1,2)"])
    sheet.append(["External formula", "='[private.xlsx]Hidden'!A1"])
    sheet["B4"] = "Visible label"
    sheet["B4"].hyperlink = "https://example.invalid/do-not-fetch"
    sheet["B5"] = "HIDDEN-LINK-SENTINEL"
    sheet["B5"].hyperlink = "https://hidden.invalid/do-not-report"
    sheet.row_dimensions[5].hidden = True
    sheet.append(["Uncached formula", "=1+1"])
    workbook.save(path)

    with zipfile.ZipFile(path) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    worksheet = parts["xl/worksheets/sheet1.xml"]
    worksheet = worksheet.replace(b"<f>SUM(1,2)</f><v />", b"<f>SUM(1,2)</f><v>3</v>")
    worksheet = worksheet.replace(
        b"<f>'[private.xlsx]Hidden'!A1</f><v />",
        b"<f>'[private.xlsx]Hidden'!A1</f><v>987654321</v>",
    )
    parts["xl/worksheets/sheet1.xml"] = worksheet
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)

    result = ingest(str(path), Settings())
    assert "SUM(1,2)" not in result.body_html
    assert ">3</td>" in result.body_html
    assert "Formula value not available" in result.body_html
    assert "private.xlsx" not in result.body_html
    assert "987654321" not in result.body_html
    assert "example.invalid" not in result.body_html
    assert "hidden.invalid" not in result.body_html
    assert "hidden.invalid" not in json.dumps(result.warnings)
    assert "External workbook value not loaded" in result.body_html
    assert _warning_codes(result) >= {
        "XLSX_EXTERNAL_REFERENCE",
        "XLSX_FORMULA_UNCALCULATED",
        "XLSX_HYPERLINK_SKIPPED",
    }
    hyperlink = next(detail for code, detail in result.warnings if code == "XLSX_HYPERLINK_SKIPPED")
    assert hyperlink.startswith("1 hyperlink")


def test_embedded_images_charts_and_hidden_anchors_are_safe(tmp_path):
    path = tmp_path / "drawings.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Drawings"
    sheet.append(["Name", "Value"])
    sheet.append(["Alpha", 3])
    sheet.append(["Beta", 5])
    sheet.add_image(SpreadsheetImage(_png()), "D2")
    sheet.add_image(SpreadsheetImage(_png("#9a3412")), "F3")
    sheet.column_dimensions["F"].hidden = True
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "H2")
    workbook.save(path)

    result = ingest(str(path), Settings())
    assert len(result.assets) == 1
    assert next(iter(result.assets)) in result.body_html
    assert 'alt="Spreadsheet image without authored alternative text."' in result.body_html
    assert "Excel chart not rendered" in result.body_html
    assert _warning_codes(result) >= {
        "XLSX_CHART_UNSUPPORTED",
        "XLSX_HIDDEN_CONTENT",
        "XLSX_IMAGE_ALT_MISSING",
    }
    hidden = next(detail for code, detail in result.warnings if code == "XLSX_HIDDEN_CONTENT")
    assert "1 drawing(s)" in hidden

    _, info = Document(str(path), str(tmp_path), Settings(), NoDiagrams(), source=result).render()
    assert sum(status == "OK" for status, _ in info["images"]) == 1


def test_merged_cells_keep_semantic_spans_without_moving_hidden_anchor_content(tmp_path):
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.freeze_panes = "A2"
    sheet.append(["First", "Second", "Hidden", "Visible"])
    sheet.merge_cells("A2:B3")
    sheet["A2"] = "Visible merged"
    sheet.merge_cells("C2:D2")
    sheet["C2"] = "HIDDEN-MERGED-SENTINEL"
    sheet.column_dimensions["C"].hidden = True
    workbook.save(path)

    result = ingest(str(path), Settings())
    assert '<td rowspan="2" colspan="2">Visible merged</td>' in result.body_html
    assert "HIDDEN-MERGED-SENTINEL" not in result.body_html


def test_linked_image_becomes_visible_warning_without_target_disclosure(tmp_path):
    path = tmp_path / "linked.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Visible"
    sheet.add_image(SpreadsheetImage(_png()), "B2")
    workbook.save(path)

    with zipfile.ZipFile(path) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    drawing_name = next(
        name for name in parts if name.startswith("xl/drawings/drawing") and name.endswith(".xml")
    )
    rels_name = "xl/drawings/_rels/" + drawing_name.rsplit("/", 1)[-1] + ".rels"
    parts[drawing_name] = parts[drawing_name].replace(b'r:embed="rId1"', b'r:link="rId1"')
    parts[rels_name] = re_sub_bytes(
        parts[rels_name],
        rb'Target="[^\"]+"',
        b'Target="https://private.invalid/image.png" TargetMode="External"',
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)

    result = ingest(str(path), Settings())
    assert not result.assets
    assert "Linked spreadsheet image not loaded" in result.body_html
    assert "private.invalid" not in result.body_html
    assert "private.invalid" not in json.dumps(result.warnings)
    assert "XLSX_LINKED_IMAGE" in _warning_codes(result)


def test_embedded_objects_are_never_opened_and_leave_a_visible_warning(tmp_path):
    path = tmp_path / "object.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Visible"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b"</worksheet>",
        b'<oleObjects><oleObject progId="Package" shapeId="1"/></oleObjects></worksheet>',
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)

    result = ingest(str(path), Settings())
    assert "Embedded spreadsheet object not loaded" in result.body_html
    assert "XLSX_OBJECT_SKIPPED" in _warning_codes(result)


def re_sub_bytes(value: bytes, pattern: bytes, replacement: bytes) -> bytes:
    import re

    changed, count = re.subn(pattern, replacement, value, count=1)
    assert count == 1
    return changed


def test_sheet_and_cell_limits_malformed_packages_and_settings_bounds(tmp_path):
    path = _workbook(tmp_path / "limits.xlsx")
    with pytest.raises(ResourceLimitError, match="sheet limit"):
        ingest(
            str(path),
            Settings(limits=replace(ResourceLimits(), max_sheet_count=2)),
        )
    with pytest.raises(ResourceLimitError, match="cell used-range limit"):
        ingest(
            str(path),
            Settings(limits=replace(ResourceLimits(), max_sheet_cells=4)),
        )
    with pytest.raises(ResourceLimitError, match="entry limit"):
        ingest(
            str(path),
            Settings(limits=replace(ResourceLimits(), max_archive_entries=1)),
        )
    with pytest.raises(ResourceLimitError, match="archive limit"):
        ingest(
            str(path),
            Settings(limits=replace(ResourceLimits(), max_archive_expanded_bytes=512)),
        )

    malformed = tmp_path / "malformed.xlsx"
    malformed.write_bytes(b"not a zip")
    with pytest.raises((InvalidOOXMLError, ResourceLimitError)):
        ingest(str(malformed), Settings())

    invalid_xml = tmp_path / "invalid-xml.xlsx"
    invalid_xml.write_bytes(path.read_bytes())
    with zipfile.ZipFile(invalid_xml) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    parts["xl/worksheets/sheet1.xml"] = b"<worksheet>"
    with zipfile.ZipFile(invalid_xml, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    with pytest.raises(InvalidOOXMLError, match="malformed OOXML part"):
        ingest(str(invalid_xml), Settings())

    with pytest.raises(ValueError, match="xlsx_sheets"):
        make_settings(xlsx_sheets="Summary")
    with pytest.raises(ValueError, match="xlsx_sheets"):
        make_settings(xlsx_sheets=("",))


def test_sixty_columns_route_through_one_custom_table_plate(tmp_path):
    path = _workbook(tmp_path / "wide.xlsx", wide=True)
    source = ingest(str(path), Settings())
    page, _ = Document(str(path), str(tmp_path), Settings(), NoDiagrams(), source=source).render()
    assert page.count('<section class="data-table table-plate') == 1
    assert "@page table-widexlsx" in page
    assert '<section class="data-table wide-table' not in page


def test_xlsx_source_theme_can_be_saved_for_reuse(tmp_path, monkeypatch):
    path = _workbook(tmp_path / "theme.xlsx")
    user = tmp_path / "themes"
    monkeypatch.setattr("platefold.theme.user_theme_dir", lambda: user)
    request = ExportRequest(
        str(path),
        str(tmp_path / "out"),
        Settings(source_theme=True),
        save_theme="workbook-theme",
    )
    saved = _save_requested_theme(request, prepare_export(request, None))
    assert saved and saved["name"] == "workbook-theme"
    payload = json.loads((user / "workbook-theme.json").read_text("utf-8"))
    assert payload["description"].endswith("XLSX by Plate|Fold.")


@pytest.mark.integration
def test_real_multisheet_wide_pdf_formats_and_hidden_absence(tmp_path, run_export):
    path = _workbook(tmp_path / "real-wide.xlsx", wide=True)
    code, output = run_export(path, "--no-toc")
    assert code == 0

    reader = PdfReader(output / "real-wide.pdf")
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    compact_text = "".join(text.split())
    sizes = {
        (round(float(page.mediabox.width)), round(float(page.mediabox.height)))
        for page in reader.pages
    }
    assert "Summary ü" in text and "Details" in text
    assert "12.50%" in compact_text and "2026-09-14" in compact_text
    assert "HIDDEN-" not in text and "Private" not in text
    assert len(sizes) >= 2

    report = json.loads((output / "_export-report.json").read_text("utf-8"))
    record = report["documents"][0]
    assert any("XLSX_HIDDEN_CONTENT" in warning for warning in record["warnings"])
    assert report["settings"]["resource_limits"]["max_sheet_cells"] == 1_000_000
