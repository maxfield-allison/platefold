"""Authored text that spells the renderer's internal placeholders stays text.

The renderer marks diagram positions, figure numbers, tile sheets and
contents page numbers in its HTML and fills them in later. A document may
contain the same characters — a page about the renderer, a code sample, an
imported cell — and must come out exactly as written, next to real diagrams,
a real contents page and tiled figures.
"""

from __future__ import annotations

import html
import json
import re
import sys

import pytest
from openpyxl import Workbook
from pypdf import PdfReader

from platefold.config import make_settings
from platefold.document import Document, finish_html

# Spellings the renderer used, or uses in some form, for its own markers.
LITERALS = [
    "@@FIGNUM@@",
    "@@FIGREF@@",
    "@@TILESET0@@",
    "@@TILESET99@@",
    "@@PG:example@@",
    "@@PG:section-1@@",
    "@@FIGURE0@@",
    "@@FIGURE1@@",
]
# Heading marker text, as authored link text.
MARKER_SPELLING = "Zq0123abcd"


class LabelledDiagrams:
    """Every fence renders to one diagram whose label is the fence's own text.

    The label passes through into the SVG as Mermaid would carry it, so a
    diagram label that spells a placeholder is tested with the rest.
    """

    def __init__(self, width: float = 400.0, height: float = 200.0):
        self.width, self.height = width, height

    def render(self, code):
        label = code.strip().splitlines()[-1].strip()
        return {
            "label": label,
            "viewBox": f"0 0 {self.width} {self.height}",
            "w": self.width,
            "h": self.height,
        }

    def svg_text(self, rec):
        return (
            '<svg id="m1" xmlns="http://www.w3.org/2000/svg" width="1" height="1" '
            'viewBox="0 0 1 1"><title>Flow</title><desc>A flow</desc>'
            f"<text>{html.escape(rec['label'])}</text></svg>"
        )


def markdown_source() -> str:
    """Every literal in prose, code, a table cell, headings, links and a diagram label."""
    inline = " ".join(f"`{literal}`" for literal in LITERALS)
    prose = " and ".join(LITERALS)
    fenced = "\n".join(LITERALS)
    cells = "\n".join(f"| {literal} | cell |" for literal in LITERALS)
    headings = "\n\n".join(f"## Section {index}\n\nBody {index}." for index in range(1, 11))
    return (
        "# Placeholder notes @@FIGREF@@\n"
        "Document: @@FIGNUM@@ notes\n"
        "Version: @@PG:example@@\n\n"
        f"Prose: {prose}.\n\n"
        f"Inline: {inline}.\n\n"
        f"```text\n{fenced}\n```\n\n"
        f"| Literal | Kind |\n|---|---|\n{cells}\n\n"
        f"A link: [{MARKER_SPELLING}](#section-1).\n\n"
        "## Heading @@PG:example@@ and @@FIGNUM@@\n\n"
        '```mermaid\nflowchart LR\n  a --> b\n  e["@@FIGNUM@@ @@TILESET0@@ @@FIGURE1@@"]\n```\n\n'
        "Between the diagrams.\n\n"
        '```mermaid\nflowchart LR\n  c --> d\n  f["@@FIGREF@@ @@PG:section-1@@"]\n```\n\n'
        f"{headings}\n"
    )


def literal_counts(text: str) -> dict[str, int]:
    return {literal: text.count(literal) for literal in [*LITERALS, MARKER_SPELLING]}


def body_of(page: str) -> str:
    return page.split("<body>", 1)[1]


@pytest.mark.parametrize("mode", ["online", "print", "tiled"])
def test_markdown_literals_survive_beside_real_figures_and_contents(tmp_path, mode):
    source = tmp_path / "notes.md"
    source.write_text(markdown_source(), encoding="utf-8")
    settings = make_settings("letter", output_mode=mode)
    # Tiled print needs a diagram larger than the sheet.
    renderer = LabelledDiagrams(3000.0, 3000.0) if mode == "tiled" else LabelledDiagrams()

    template, info = Document(str(source), str(tmp_path), settings, renderer).render()
    pages = {anchor: index + 2 for index, anchor in enumerate(info["order"])}
    measuring = finish_html(template, markers=True)
    final = finish_html(template, toc_pages=pages)

    authored = markdown_source()
    for page in (measuring, final):
        counts = literal_counts(body_of(page))
        # Each literal survives everywhere it was written: the body, the
        # heading and its contents entry, the metadata and the diagram labels.
        for literal, count in counts.items():
            assert count >= authored.count(literal), (literal, count)
        assert "IndexError" not in page
    # Exactly the two real diagrams are numbered, in reading order.
    assert info["figures"] == 2
    assert re.findall(r"<figcaption>Figure (\d+)", final) == ["1", "2"]
    # The real contents page gets real numbers from the measuring pass.
    assert info["has_toc"]
    assert 'href="#section-1"><span class="t">Section 1</span>' in final
    assert re.search(
        r'href="#section-1"><span class="t">Section 1</span><span class="dots"></span>'
        rf'<span class="pg">{pages["section-1"]}</span>',
        final,
    )
    # The heading markers' text is removed from the delivered pass only; the
    # authored link text with the same shape is kept in both.
    assert f'href="#section-1">{MARKER_SPELLING}</a>' in final
    assert f'href="#section-1">{MARKER_SPELLING}</a>' in measuring
    if mode == "tiled":
        tiled = info["tiled_figures"]
        assert [entry["figure"] for entry in tiled] == [1, 2]
        assert final.count('<span class="tile-name">Figure 1') == tiled[0]["sheets"]
        assert final.count('<span class="tile-name">Figure 2') == tiled[1]["sheets"]
    else:
        assert info["tiled_figures"] == []


def test_placeholder_spelling_without_figures_is_left_alone(tmp_path):
    """The reviewed reproduction: numbering and a tile lookup ran on prose."""
    source = tmp_path / "tokens.md"
    source.write_text("Literal `@@FIGNUM@@` and `@@PG:example@@`.", encoding="utf-8")
    template, info = Document(str(source), str(tmp_path), make_settings(), None).render()
    assert "<code>@@FIGNUM@@</code> and <code>@@PG:example@@</code>" in finish_html(template)
    assert info["figures"] == 0

    source.write_text("Literal `@@TILESET99@@`.", encoding="utf-8")
    template, info = Document(str(source), str(tmp_path), make_settings(), None).render()
    assert "<code>@@TILESET99@@</code>" in finish_html(template)
    assert info["tiled_figures"] == []


def test_copying_a_placeholder_spelling_into_the_document_changes_it(tmp_path):
    """The spelling is derived from the document, so it cannot be planted in it."""
    source = tmp_path / "flow.md"
    text = "```mermaid\nflowchart LR\n  a --> b\nlabel\n```\n"
    source.write_text(text, encoding="utf-8")
    first, _ = Document(str(source), str(tmp_path), make_settings(), LabelledDiagrams()).render()
    again, _ = Document(str(source), str(tmp_path), make_settings(), LabelledDiagrams()).render()
    # The same input renders the same way.
    assert first == again
    planted = f"{first.placeholders.fignum} {first.placeholders.page('x')}"
    source.write_text(text + f"\nPlanted `{planted}`.\n", encoding="utf-8")
    second, info = Document(
        str(source), str(tmp_path), make_settings(), LabelledDiagrams()
    ).render()
    assert second.placeholders.key != first.placeholders.key
    assert f"<code>{planted}</code>" in finish_html(second)
    assert info["figures"] == 1


@pytest.mark.skipif(
    sys.platform == "win32", reason="Windows file names cannot hold the ':' every placeholder has"
)
def test_a_file_name_spelling_a_placeholder_keeps_it_as_the_title(tmp_path):
    """An untitled document is titled by its file name, which is authored text too."""
    text = "No title heading here.\n\n## Section\n\nBody.\n"
    seed = tmp_path / "seed.md"
    seed.write_text(text, encoding="utf-8")
    first, _ = Document(str(seed), str(tmp_path), make_settings(toc=True), None).render()
    # The same text in a file named with the spelling that text would get.
    name = f"report {first.placeholders.page('x')}.md"
    named = tmp_path / name
    named.write_text(text, encoding="utf-8")
    template, _ = Document(str(named), str(tmp_path), make_settings(toc=True), None).render()
    assert f"<title>{html.escape(name)}</title>" in finish_html(template)


def _render_imported(tmp_path, path) -> str:
    template, _ = Document(str(path), str(tmp_path), make_settings(), None).render()
    return finish_html(template)


def test_csv_cells_keep_placeholder_spelling(tmp_path):
    path = tmp_path / "cells.csv"
    path.write_text("Value\n" + "\n".join(LITERALS) + "\n", encoding="utf-8")
    page = _render_imported(tmp_path, path)
    for literal in LITERALS:
        assert f"<td>{literal}</td>" in page


def test_xlsx_cells_keep_placeholder_spelling(tmp_path):
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cells"
    sheet.append(["Value"])
    for literal in LITERALS:
        sheet.append([literal])
    path = tmp_path / "cells.xlsx"
    book.save(path)
    page = _render_imported(tmp_path, path)
    for literal in LITERALS:
        assert f"{literal}</td>" in page


def test_docx_text_and_cells_keep_placeholder_spelling_beside_a_real_diagram(tmp_path):
    from test_docx_ingest import make_docx

    def paragraph(text: str) -> str:
        return f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'

    cells = "".join(f"<w:tr><w:tc>{paragraph(literal)}</w:tc></w:tr>" for literal in LITERALS)
    body = (
        "".join(paragraph(f"Prose {literal}.") for literal in LITERALS)
        + f"<w:tbl>{cells}</w:tbl>"
        + '<w:p><w:pPr><w:pStyle w:val="CodeBlock"/></w:pPr><w:r>'
        '<w:t xml:space="preserve">```mermaid\nflowchart LR\n A --&gt; B\n```</w:t></w:r></w:p>'
    )
    path = make_docx(tmp_path / "cells.docx", body=body)
    template, info = Document(
        str(path), str(tmp_path), make_settings(), LabelledDiagrams()
    ).render()
    page = finish_html(template)
    assert info["figures"] == 1
    for literal in LITERALS:
        assert f"Prose {literal}." in page
        assert page.count(literal) >= 2


@pytest.mark.integration
def test_real_export_keeps_literals_beside_mermaid_contents_and_tiles(tmp_path, run_export):
    source = tmp_path / "notes.md"
    source.write_text(markdown_source(), encoding="utf-8")
    for flags in ([], ["--output-mode", "tiled"]):
        code, out = run_export(source, *flags)
        assert code == 0
        text = "".join(
            "".join((page.extract_text() or "").split())
            for page in PdfReader(out / "notes.pdf").pages
        )
        for literal in [*LITERALS, MARKER_SPELLING]:
            assert literal in text, (flags, literal)
        report = json.loads((out / "_export-report.json").read_text(encoding="utf-8"))
        assert report["documents"][0]["figures"] == 2
        assert report["documents"][0]["diagram_errors"] == []
