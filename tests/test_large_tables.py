"""Large tables print in parts, and the joined PDF matches one pass.

A document with more table cells than one browser pass may lay out is printed
in parts and joined. These tests print documents both ways, with a small
per-pass budget so the parts are many, and compare every page image, the
whole tagged structure and the export report. One more exports a generated
million-cell workbook at the default limits. Everything is generated at run
time.
"""

import csv
import hashlib
import os
import random
import re
import shutil
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook
from pypdf import PdfReader
from test_pdf_ua_fixtures import grid_problems, row_widths, table_cells

from platefold.config import Settings, make_settings
from platefold.ingest import ingest
from platefold.limits import ResourceLimitError, ResourceLimits
from platefold.pdfobjects import PdfFile, Ref
from platefold.service import ExportRequest, export
from platefold.style import build_css

LIMITS = ResourceLimits()
COLUMNS = 20


def _values(row: int, rng: random.Random, columns: int = COLUMNS) -> list:
    """One row of mixed text and numbers, like an ordinary data export."""
    return [
        None  # an empty cell, which the PDF must still tag
        if (row + column) % 13 == 0
        else f"item {row}-{column} {rng.choice(['north', 'south', 'east', 'west'])}"
        if column % 3 == 0
        else round(rng.random() * 100_000, 2)
        for column in range(columns)
    ]


def _workbook(path: Path, sheets: int, rows: int, columns: int = COLUMNS) -> Path:
    """Sheets of ``rows`` rows each, the header row included."""
    rng = random.Random(7)
    book = Workbook(write_only=True)
    for index in range(1, sheets + 1):
        sheet = book.create_sheet(f"Data {index}")
        sheet.append([f"Column {column}" for column in range(1, columns + 1)])
        for row in range(rows - 1):
            sheet.append(_values(row, rng, columns))
    book.save(path)
    return path


def _corners(path: Path, sheets: int, rows: int) -> Path:
    """Sheets whose used range spans ``rows`` x COLUMNS but hold two cells."""
    book = Workbook()
    for index in range(1, sheets + 1):
        sheet = book.active if index == 1 else book.create_sheet()
        sheet.title = f"Data {index}"
        sheet["A1"] = "first"
        sheet.cell(row=rows, column=COLUMNS, value="last")
    book.save(path)
    return path


def _markdown(path: Path) -> Path:
    """A contents page, lists, links across parts, a long table and a short one.

    List items carry emphasis and some cells are empty, including a whole row,
    so the joined PDF exercises role mapping and empty-cell tagging for the
    PDF/UA evidence.
    """
    rng = random.Random(3)
    lines = ["# Report", "", "See [the notes](#notes) and [the web](https://example.com).", ""]
    lines += [
        f"## Section {i}\n\n- item with **bold**\n- item with *emphasis*\n" for i in range(1, 10)
    ]
    lines += ["## Data", "", "| Region | Item | Value | Note |", "| --- | --- | --- | --- |"]
    for i in range(700):
        words = " ".join(rng.choice(["alpha", "beta", "gamma"]) for _ in range(rng.randint(1, 12)))
        note = "" if i % 9 == 0 else f"**{words}**" if i % 11 == 0 else words
        if i == 350:
            lines.append("|  |  |  |  |")
        else:
            lines.append(f"| R{i % 7} | item {i} | {rng.random() * 1000:.2f} | {note} |")
    lines += ["", "## Notes", "", "Notes after the table. " * 30, "", "| A | B |", "| --- | --- |"]
    lines += [f"| a{i} | {'text ' * (i % 5)} |" for i in range(40)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _headerless_csv(path: Path) -> Path:
    rng = random.Random(1)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in range(400):
            writer.writerow(
                [
                    "" if (row + c) % 11 == 0 else f"v{row}-{c}-{'x' * rng.randint(1, 20)}"
                    for c in range(12)
                ]
            )
    return path


def _images(path: Path) -> list[str]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        return [
            hashlib.sha256(document[i].render(scale=1).to_pil().tobytes()).hexdigest()
            for i in range(len(document))
        ]
    finally:
        document.close()


def _structure(path: Path) -> dict:
    """The tagged-PDF record, with element IDs replaced by their order of use."""
    from tools.a11y.inspect_pdf import inspect

    record = {k: v for k, v in inspect(path).items() if k not in ("sha256", "bytes", "path")}
    ids: dict[str, str] = {}
    for node in record["nodes"]:
        if "/ID" in node:
            ids.setdefault(node["/ID"], f"id{len(ids)}")

    def rename(value):
        if isinstance(value, dict):
            return {key: rename(item) for key, item in value.items()}
        if isinstance(value, list):
            return [rename(item) for item in value]
        return ids.get(value, value) if isinstance(value, str) else value

    return rename(record)


def _export(source: Path, out: Path, per_pass: int, **options) -> dict:
    limits = replace(ResourceLimits(), max_cells_per_pass=per_pass)
    settings = make_settings(document_language="en", limits=limits, **options)
    result = export(ExportRequest(str(source), str(out), settings))
    assert result.exit_code == 0, result.report
    return result.report["documents"][0]


def _keep_evidence(pdf: Path, name: str) -> None:
    """Copy a joined PDF where CI validates it against PDF/UA-1 with veraPDF."""
    directory = os.environ.get("PLATEFOLD_PDF_EVIDENCE_DIR")
    if directory:
        Path(directory).mkdir(parents=True, exist_ok=True)
        shutil.copyfile(pdf, Path(directory) / f"{name}.pdf")


def test_the_stylesheet_has_no_relational_selectors():
    # Chromium re-evaluated a :has() table rule for every cell each time the
    # parser added a row, which made loading and printing quadratic.
    rules = re.sub(r"/\*.*?\*/", "", build_css({}, Settings()), flags=re.S)
    assert ":has(" not in rules


def test_the_table_limits_bound_work_and_the_pass_budget_bounds_memory():
    assert LIMITS.max_workbook_cells == LIMITS.max_sheet_cells == LIMITS.max_csv_cells
    assert LIMITS.max_document_cells == 1_000_000
    assert LIMITS.max_cells_per_pass < LIMITS.max_document_cells


def test_a_workbook_one_row_past_the_limit_is_refused_before_it_is_read(tmp_path):
    rows = LIMITS.max_workbook_cells // 2 // COLUMNS + 1
    path = _corners(tmp_path / "wide.xlsx", 2, rows)
    with pytest.raises(ResourceLimitError, match="combined used-range limit; select fewer"):
        ingest(str(path), Settings())


@pytest.mark.integration
@pytest.mark.parametrize("kind", ["markdown", "headerless-csv-print", "workbook"])
def test_a_document_printed_in_parts_matches_one_pass(tmp_path, kind):
    if kind == "markdown":
        source, options = _markdown(tmp_path / "report.md"), {}
    elif kind == "workbook":
        source, options = _workbook(tmp_path / "sheets.xlsx", 10, 150, 8), {}
    else:
        source = _headerless_csv(tmp_path / "plain.csv")
        options = {"output_mode": "print", "csv_header": False}

    parts = _export(source, tmp_path / "parts", 600, **options)
    whole = _export(source, tmp_path / "whole", 10**9, **options)

    split = next((tmp_path / "parts").glob("*.pdf"))
    single = next((tmp_path / "whole").glob("*.pdf"))
    assert parts == whole  # pages, warnings and font evidence alike
    assert _images(split) == _images(single)
    assert _structure(split) == _structure(single)
    # Every row keeps all its cells, empty ones included (PDF/UA-1 7.2-43).
    with PdfReader(split) as reader:
        tables = row_widths(reader)
    assert tables and all(len(set(rows)) == 1 for rows in tables)
    _keep_evidence(split, f"parts-{kind}")


@pytest.mark.integration
def test_a_million_cell_workbook_exports_in_parts(tmp_path):
    source = _workbook(tmp_path / "million.xlsx", 4, LIMITS.max_sheet_cells // 4 // COLUMNS)

    document = _export(source, tmp_path / "out", LIMITS.max_cells_per_pass)

    pdf_path = tmp_path / "out" / "million.pdf"
    assert not document["warnings"]
    pdf = PdfFile(pdf_path)
    pages = pdf.pages()
    assert document["pages"] == len(pages) > 1500
    # One structure-parent key per page, numbered as one pass numbers them,
    # and every reference in every object resolves.
    keys = [pdf.get(Ref(num))[b"/StructParents"] for num in pages]
    assert keys == list(range(len(pages)))
    tree = pdf.get(pdf.root[b"/StructTreeRoot"])
    assert sorted(key for key, _value in pdf.number_tree(tree[b"/ParentTree"])) == keys
    reference = re.compile(rb"(?<![\w.])(\d+) 0 R(?!\w)")
    for num in pdf.order:
        missing = [n for n in reference.findall(pdf.raw(num).head) if int(n) not in pdf.offsets]
        assert not missing, (num, missing)
    with PdfReader(pdf_path) as reader:
        assert [item.title for item in reader.outline] == [f"Data {i}" for i in range(1, 5)]
        assert "item 12498-18" in reader.pages[-1].extract_text()


# ------------------------------------------------------- authored strings

# Alternative text and links that read like PDF syntax. The first image prints
# in the first part, the last in a part that continues the table.
_STRINGS_MARKDOWN = """# Strings

![Literal 1 0 R](shape.svg)

![Escaped \\) and \\( around 2 0 R](shape.svg)

![Nested (3 0 R (4 0 R)) and 50% 5 0 R](shape.svg)

See [the notes 6 0 R](#notes) and [the web](https://example.com/a(1%200%20R)?b=7+0+R).

| Region | Value |
| --- | --- |
{rows}

## Notes

![node00000004](shape.svg)

![After the table 8 0 R /StructParent 1](shape.svg)

Back to [the strings 9 0 R](#strings).
"""
_ALTERNATIVES = [
    "Literal 1 0 R",
    "Escaped ) and ( around 2 0 R",
    "Nested (3 0 R (4 0 R)) and 50% 5 0 R",
    "node00000004",
    "After the table 8 0 R /StructParent 1",
]


def _dictionary_strings(path: Path) -> list[tuple[bytes, bytes]]:
    """Every string of the fields that carry document text, in every object, by field.

    Each part embeds its own fonts, so their strings are left out.
    """
    from platefold.pdfobjects import scan, string_bytes

    fields = {b"/Alt", b"/ActualText", b"/E", b"/Lang", b"/T", b"/Contents", b"/URI", b"/D"}
    fields |= {b"/Dest", b"/Title"}
    pdf = PdfFile(path)
    found = []
    for num in pdf.order:
        head = pdf.raw(num).head
        syntax = scan(head)
        for start, end in syntax.strings:
            field = syntax.field(start)
            if field in fields:
                found.append((field, string_bytes(head[start:end])))
    return sorted(found)


@pytest.mark.integration
def test_strings_that_read_like_pdf_syntax_print_the_same_in_parts(tmp_path):
    (tmp_path / "shape.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20">'
        '<rect width="20" height="20" fill="#357"/></svg>',
        encoding="utf-8",
    )
    source = tmp_path / "strings.md"
    rows = "\n".join(f"| R{i} | value {i} |" for i in range(300))
    source.write_text(_STRINGS_MARKDOWN.replace("{rows}", rows), encoding="utf-8")

    parts = _export(source, tmp_path / "parts", 100)
    whole = _export(source, tmp_path / "whole", 10**9)

    split, single = tmp_path / "parts" / "strings.pdf", tmp_path / "whole" / "strings.pdf"
    assert parts == whole
    assert _structure(split)["alternatives"] == _ALTERNATIVES
    assert _structure(split) == _structure(single)
    assert _dictionary_strings(split) == _dictionary_strings(single)
    assert _images(split) == _images(single)
    _keep_evidence(split, "parts-strings")


# ------------------------------------------------------ cells spanning rows


def _spanned_table(count: int, spans: dict[tuple[int, int], int], columns: int) -> str:
    """A raw HTML table whose cell at (row, column) spans ``spans[(row, column)]`` rows."""
    covered: dict[int, int] = {}
    body = []
    for r in range(count):
        cells = []
        for c in range(columns):
            if covered.get(c, 0) > r:
                continue
            span = spans.get((r, c), 1)
            covered[c] = r + span
            rowspan = f' rowspan="{span}"' if span > 1 else ""
            cells.append(f"<td{rowspan}>{'g' if span > 1 else f'row {r} col {c}'}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    head = "".join(f"<th>Column {c}</th>" for c in range(columns))
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _export_counting_passes(source: Path, out: Path, per_pass: int, passes: list[int]):
    """Export, recording in ``passes`` the table cells of every pass the browser is given."""
    from platefold import chunked, htmlplan

    original = chunked.part_html

    def recording(*args, **kwargs):
        text, stand_ins = original(*args, **kwargs)
        passes.append(chunked.table_cells(htmlplan.scan(text)))
        return text, stand_ins

    limits = replace(ResourceLimits(), max_cells_per_pass=per_pass)
    settings = make_settings(document_language="en", limits=limits)
    chunked.part_html = recording
    try:
        return export(ExportRequest(str(source), str(out), settings))
    finally:
        chunked.part_html = original


@pytest.mark.integration
def test_rows_joined_by_a_long_span_are_refused_before_a_pass_holds_them(tmp_path):
    source = tmp_path / "joined.md"
    table = _spanned_table(300, {(0, 0): 300}, 2)
    source.write_text(f"# Joined\n\n{table}\n", encoding="utf-8")

    passes: list[int] = []
    with pytest.raises(ResourceLimitError, match="cells spanning rows"):
        _export_counting_passes(source, tmp_path / "out", 50, passes)
    assert passes == []


@pytest.mark.integration
def test_spanning_cells_print_in_parts_within_the_limit(tmp_path):
    # Chained spans join rows 0-61; after them ordinary merged cells every ten rows.
    spans = {(r, (r // 2) % 2): 4 for r in range(0, 60, 2)}
    spans |= {(r, 0): 2 for r in range(70, 300, 10)}
    source = tmp_path / "spans.md"
    table = _spanned_table(300, spans, 3)
    source.write_text(f"# Spans\n\n{table}\n", encoding="utf-8")

    passes: list[int] = []
    result = _export_counting_passes(source, tmp_path / "parts", 150, passes)
    assert result.exit_code == 0, result.report
    whole = _export(source, tmp_path / "whole", 10**9)

    assert len(passes) > 2 and max(passes) <= 4 * 150
    split, single = tmp_path / "parts" / "spans.pdf", tmp_path / "whole" / "spans.pdf"
    assert result.report["documents"][0] == whole
    assert _structure(split) == _structure(single)
    # Column widths included: the hidden rows carry rows with and under spans.
    assert _images(split) == _images(single)


# Words of different widths, so that which cell is widest decides the columns.
_WORDS = ["alpha", "iii", "WWW", "Supercalifragilistic", "mm", "north by northwest", "x", "gamma"]


def _merged_table(
    count: int, columns: int, cells: dict[tuple[int, int], tuple[int, int]], head: str = ""
) -> str:
    """A raw HTML table whose cell at (row, column) spans ``cells[(row, column)]`` rows, columns.

    Each cell holds a few words; a row every cell of which is covered by a
    cell above it has no cells at all. A cell spans no further across than
    the columns free in its row, so no two cells overlap, which HTML calls a
    table model error; a span may run past the last row, where layout ends it.
    """
    rng = random.Random(count * columns + len(cells))
    taken: set[tuple[int, int]] = set()
    body = []
    for r in range(count):
        row = []
        for c in range(columns):
            if (r, c) in taken:
                continue
            down, wanted = cells.get((r, c), (1, 1))
            across = 1
            while across < min(wanted, columns - c) and (r, c + across) not in taken:
                across += 1
            taken |= {(r + i, c + j) for i in range(down) for j in range(across)}
            spans = (f' rowspan="{down}"' if down > 1 else "") + (
                f' colspan="{across}"' if across > 1 else ""
            )
            text = " ".join(rng.choice(_WORDS) for _ in range(rng.randint(1, 4 + 4 * across)))
            row.append(f"<td{spans}>{text}</td>")
        body.append("<tr>" + "".join(row) + "</tr>")
    head = head or "<tr>" + "".join(f"<th>Column {c}</th>" for c in range(columns)) + "</tr>"
    return f"<table><thead>{head}</thead><tbody>{''.join(body)}</tbody></table>"


def _merged_cases() -> dict[str, tuple[str, int]]:
    """Tables with merged cells, and the per-pass budget that splits each into many parts."""
    boundary = {(r, 0): (2, 1) for r in range(0, 300, 2)}  # every part starts under a span
    chained = {(r, (r // 2) % 2): (4, 1) for r in range(0, 120, 2)}
    chained |= {(r, 0): (3, 1) for r in range(130, 300, 7)}  # the last runs past the end
    both = {}
    for r in range(0, 300, 6):
        both[(r, 0)] = (3, 2)  # rows and columns at once
        both[(r + 3, 1)] = (2, 3)
    both |= {(r + 5, 0): (1, 4) for r in range(0, 300, 25)}  # a section row across the table
    both |= {(r, 0): (3, 4) for r in (101, 203)}  # rows with no cells of their own
    header = {(r, 1): (2, 2) for r in range(0, 300, 5)} | {(r, 0): (3, 1) for r in range(1, 300, 9)}
    # A cell over three rows, the middle one without cells of its own, the
    # third holding the other columns: that row is still under the first cell.
    partly = {}
    for r in range(0, 300, 3):
        partly |= {(r, 0): (3, 1), (r, 1): (2, 1), (r, 2): (2, 1)}
    grouped = (
        '<tr><th rowspan="2">Item</th><th colspan="2">First group</th>'
        '<th colspan="2">Second group of columns</th></tr>'
        "<tr><th>A</th><th>Bee</th><th>C</th><th>Dee</th></tr>"
    )
    return {
        "rowspan-at-part-boundary": (_merged_table(300, 3, boundary), 150),
        "chained-overlapping-spans": (_merged_table(300, 3, chained), 150),
        "rowspans-and-colspans": (_merged_table(300, 4, both), 150),
        "spanning-header": (_merged_table(300, 5, header, grouped), 150),
        "partly-covered-rows": (_merged_table(300, 3, partly), 150),
    }


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
@pytest.mark.parametrize("case", sorted(_merged_cases()))
def test_merged_cells_print_the_same_in_parts(tmp_path, case):
    table, budget = _merged_cases()[case]
    source = tmp_path / "merged.md"
    source.write_text(f"# Merged\n\nBefore.\n\n{table}\n\nAfter the table.\n", encoding="utf-8")

    parts = _export(source, tmp_path / "parts", budget)
    whole = _export(source, tmp_path / "whole", 10**9)

    split, single = tmp_path / "parts" / "merged.pdf", tmp_path / "whole" / "merged.pdf"
    assert parts == whole
    assert _images(split) == _images(single)
    assert _structure(split) == _structure(single)
    _assert_one_grid(split)
    _keep_evidence(split, f"parts-{case}")


def _assert_one_grid(path: Path) -> None:
    """The tagged cells of every table form one grid, as PDF/UA-1 requires."""
    tables = table_cells(PdfReader(path))
    assert tables
    for rows in tables:
        assert grid_problems(rows) == []


def _merged_workbook(path: Path) -> Path:
    """A sheet of merged ranges: down one column, across a row, blocks, a band across all."""
    rng = random.Random(11)
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append([f"Column {column}" for column in range(1, 7)])
    for _row in range(2, 402):
        sheet.append(
            [" ".join(rng.choice(_WORDS) for _ in range(rng.randint(1, 5))) for _ in range(6)]
        )
    for row in range(2, 402, 7):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=1)
    for row in range(5, 402, 13):
        sheet.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=4)
    for row in range(10, 402, 50):
        sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
    sheet.merge_cells(start_row=180, start_column=1, end_row=182, end_column=6)
    book.save(path)
    return path


def _partly_merged_workbook(path: Path) -> Path:
    """A column merged over three rows beside two merged over the first two of them."""
    rng = random.Random(5)
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Group", "Item", "Note"])
    for _row in range(120):
        sheet.append(
            [" ".join(rng.choice(_WORDS) for _ in range(rng.randint(1, 6))) for _ in range(3)]
        )
    for row in range(2, 120, 3):
        sheet.merge_cells(f"A{row}:A{row + 2}")
        sheet.merge_cells(f"B{row}:B{row + 1}")
        sheet.merge_cells(f"C{row}:C{row + 1}")
    book.save(path)
    return path


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
@pytest.mark.parametrize(
    "name,build,budget",
    [
        ("merged-workbook", _merged_workbook, 400),
        ("partly-merged-workbook", _partly_merged_workbook, 60),
    ],
)
def test_a_workbook_with_merged_cells_prints_the_same_in_parts(tmp_path, name, build, budget):
    source = build(tmp_path / "merged.xlsx")

    parts = _export(source, tmp_path / "parts", budget)
    whole = _export(source, tmp_path / "whole", 10**9)

    split = next((tmp_path / "parts").glob("*.pdf"))
    single = next((tmp_path / "whole").glob("*.pdf"))
    assert parts == whole
    assert _images(split) == _images(single)
    assert _structure(split) == _structure(single)
    _assert_one_grid(split)
    _keep_evidence(split, f"parts-{name}")
