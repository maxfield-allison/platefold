"""Documents that hold the structures PDF/UA checks are sensitive to.

Inline emphasis becomes PDF 2.0 structure types that PDF/UA-1 requires to be
role-mapped, an empty table cell used to drop out of its row, and a merged
cell used to be tagged over rows its table does not hold. Each export here is
also kept as veraPDF evidence for CI (PLATEFOLD_PDF_EVIDENCE_DIR).
"""

import csv
import hashlib
import json
import os
import shutil
import time
from pathlib import Path

import pytest
from openpyxl import Workbook
from pypdf import PdfReader
from pypdf.generic import DictionaryObject
from test_docx_ingest import indented_row, make_docx, merge_table

from platefold import table_spans
from platefold.accessibility import fill_empty_cells
from platefold.config import make_settings
from platefold.ingest import docx
from platefold.pdf_semantics import STANDARD_TYPES, role_map
from platefold.service import ExportRequest, export
from platefold.theme import BUILTIN_DIR, discover_themes, is_builtin

EMPHASIS = """# Emphasis

Plain **bold**, *italic*, ***both***, `code`, ~~struck~~, <sup>sup</sup>, <sub>sub</sub>,
<mark>mark</mark>, <kbd>Ctrl</kbd>, <abbr title="HyperText">HTML</abbr>, <q>quote</q>,
<small>small</small>, <u>under</u>, <cite>cite</cite>, <dfn>term</dfn>, <var>x</var>,
<samp>out</samp>, <ins>ins</ins>, <del>del</del>, <b>b</b>, <i>i</i>, <s>s</s> and a
**[bold link](https://example.com)**.

- first item with **bold**
- second with *italic* and `code`
  - nested **bold** item

1. ordered **bold**
2. ordered *italic*

> quoted **bold** and *italic*

## Heading with **bold** and *italic*

| Header **bold** | *Italic* header |
| --- | --- |
| cell **bold** | cell *italic* |

<dl><dt>Term <b>bold</b></dt><dd>Definition <em>em</em></dd></dl>
"""

EMPTY_MARKDOWN = """# Empty cells

| A | B | C |
| --- | --- | --- |
| a1 | text | c1 |
| a2 |  | c2 |
|  |  |  |
| a4 | **b4** |  |

|  | Header |
| --- | --- |
| row | value |

<table><thead><tr><th scope="col">H1</th><th scope="col"></th></tr></thead>
<tbody><tr><td colspan="2">spanning</td></tr><tr><td><br></td><td>x</td></tr></tbody></table>
"""

DOCX_BODY = (
    '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Bold run</w:t></w:r><w:r><w:t xml:space="preserve"> and </w:t></w:r>'
    "<w:r><w:rPr><w:i/></w:rPr><w:t>italic run</w:t></w:r></w:p>"
    "<w:tbl><w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:p><w:r><w:t>H1</w:t></w:r></w:p></w:tc>"
    "<w:tc><w:p/></w:tc><w:tc><w:p><w:r><w:t>H3</w:t></w:r></w:p></w:tc></w:tr>"
    "<w:tr><w:tc><w:p/></w:tc><w:tc><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>bold cell</w:t></w:r></w:p></w:tc>"
    "<w:tc><w:p/></w:tc></w:tr>"
    '<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>merged</w:t></w:r></w:p></w:tc>'
    "<w:tc><w:p/></w:tc></w:tr></w:tbl>"
)


def _source(kind: str, root: Path) -> tuple[Path, dict]:
    if kind == "emphasis":
        path = root / "emphasis.md"
        path.write_text(EMPHASIS, encoding="utf-8")
        return path, {}
    if kind in {"empty-markdown", "empty-markdown-print"}:
        path = root / "empty.md"
        path.write_text(EMPTY_MARKDOWN, encoding="utf-8")
        return path, {"output_mode": "print"} if kind.endswith("print") else {}
    if kind in {"empty-csv", "empty-csv-headerless"}:
        path = root / "empty.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows(
                [["Name", "", "Note"], ["alpha", "1", ""], ["", "", ""], ["gamma", "3", "x"]]
            )
        return path, {"csv_header": kind == "empty-csv"}
    if kind == "empty-xlsx":
        path = root / "empty.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Gaps"
        for row in (["Name", None, "Note"], [None, None, None], ["a", 2, None], ["b", None, "end"]):
            sheet.append(row)
        book.save(path)
        return path, {}
    if kind == "docx":
        return make_docx(root / "emphasis.docx", body=DOCX_BODY), {}
    raise AssertionError(kind)


def _elements(node) -> list[DictionaryObject]:
    """A structure element's child elements (not marked-content references)."""
    kids = node.get("/K", [])
    kids = kids if isinstance(kids, list) else [kids]
    children = []
    for kid in kids:
        kid = kid.get_object()
        if isinstance(kid, DictionaryObject) and kid.get("/Type", "/StructElem") == "/StructElem":
            children.append(kid)
    return children


def _table_span(cell: DictionaryObject, key: str) -> int:
    attributes = cell.get("/A")
    attributes = attributes.get_object() if attributes is not None else []
    for owner in attributes if isinstance(attributes, list) else [attributes]:
        owner = owner.get_object()
        if isinstance(owner, DictionaryObject) and key in owner:
            return int(owner[key])
    return 1


def table_cells(reader: PdfReader) -> list[list[list[tuple[int, int]]]]:
    """For each table, each row's cells as their tagged (row span, column span)."""
    tables, pending = [], [reader.root_object["/StructTreeRoot"].get_object()]
    while pending:
        node = pending.pop(0)
        if node.get("/S") == "/Table":
            rows, groups = [], _elements(node)
            while groups:
                item = groups.pop(0)
                if item.get("/S") == "/TR":
                    cells = [c for c in _elements(item) if c.get("/S") in ("/TD", "/TH")]
                    rows.append(
                        [(_table_span(c, "/RowSpan"), _table_span(c, "/ColSpan")) for c in cells]
                    )
                else:
                    groups[:0] = _elements(item)
            tables.append(rows)
        pending[:0] = _elements(node)
    return tables


def row_widths(reader: PdfReader) -> list[list[int]]:
    """For each table, the number of columns each row spans."""
    return [[sum(across for _, across in row) for row in rows] for rows in table_cells(reader)]


def grid_problems(rows: list[list[tuple[int, int]]]) -> list[str]:
    """How a tagged table's cells fail to form one grid, as PDF/UA-1 7.2-15, 7.2-42 and 7.2-43 check.

    Each cell takes the next column no cell above covers. A cell must not
    cover a place another cell holds or reach past the last row, and every
    row must fill the same columns, counting cells from rows above.
    """
    problems = []
    held: dict[tuple[int, int], int] = {}
    for r, cells in enumerate(rows):
        column = 0
        for down, across in cells:
            while (r, column) in held:
                column += 1
            if not 1 <= down <= len(rows) - r:
                problems.append(f"row {r} column {column}: spans {down} of {len(rows) - r} rows")
            for place in ((r + i, column + j) for i in range(max(down, 1)) for j in range(across)):
                if place in held:
                    problems.append(f"row {r} column {column}: overlaps a cell at {place}")
                held[place] = r
            column += across
    widths = [sorted(c for (r, c) in held if r == row) for row in range(len(rows))]
    if any(width != widths[0] or width != list(range(len(width))) for width in widths):
        problems.append(f"rows fill different columns: {sorted({len(w) for w in widths})}")
    return problems


def _roles(reader: PdfReader) -> set[str]:
    roles, pending = set(), [reader.root_object["/StructTreeRoot"].get_object()]
    while pending:
        node = pending.pop()
        if "/S" in node:
            roles.add(str(node["/S"]))
        pending += _elements(node)
    return roles


@pytest.mark.integration
@pytest.mark.parametrize(
    "kind",
    [
        "emphasis",
        "empty-markdown",
        "empty-markdown-print",
        "empty-csv",
        "empty-csv-headerless",
        "empty-xlsx",
        "docx",
    ],
)
def test_emphasis_is_role_mapped_and_every_row_keeps_its_cells(tmp_path, kind):
    source, options = _source(kind, tmp_path)
    settings = make_settings(document_language="en", **options)
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        tree = reader.root_object["/StructTreeRoot"]
        mapping = tree.get("/RoleMap", {})
        mapping = {str(k): str(v) for k, v in mapping.get_object().items()} if mapping else {}
        for role in _roles(reader) - STANDARD_TYPES:
            assert mapping.get(role) in STANDARD_TYPES, f"{role} is not role-mapped"
        tables = row_widths(reader)
        if kind != "emphasis":
            assert tables
        for rows in tables:
            assert len(set(rows)) == 1, rows
    _keep_evidence(final, f"fixture-{kind}")


def _keep_evidence(pdf: Path, name: str) -> None:
    directory = os.environ.get("PLATEFOLD_PDF_EVIDENCE_DIR")
    if directory:
        Path(directory).mkdir(parents=True, exist_ok=True)
        shutil.copyfile(pdf, Path(directory) / f"{name}.pdf")


# Merged cells whose markup spans rows the tagged table does not hold: rows
# with no cells of their own, spans past the end of their row group, a span of
# 0, a footer written before the body, and end tags left for the parser.
MERGED_MARKDOWN = """# Merged cells

<table><thead><tr><th>A</th><th>B</th><th>C</th></tr></thead><tbody>
<tr><td rowspan="3" colspan="3">a block over three rows</td></tr><tr></tr><tr></tr>
<tr><td rowspan="2">left</td><td>b</td><td>c</td></tr><tr><td>d</td><td>e</td></tr>
<tr><td rowspan="2">under</td><td colspan="2" rowspan="2">wide</td></tr><tr></tr>
<tr><td>f</td><td>g</td><td rowspan="4">past the end</td></tr>
</tbody><tbody><tr><td>h</td><td>i</td><td>j</td></tr></tbody></table>

<table><thead><tr><th rowspan="3">Item</th><th>Group</th></tr><tr><th>Part</th></tr></thead>
<tfoot><tr><td>total</td><td>9</td></tr></tfoot>
<tbody><tr><td rowspan="0">to the end</td><td>1</td></tr><tr><td>2</td></tr><tr></tr>
<tr><td>3</td></tr></tbody></table>

<table><tr><th>X<th>Y<tr><td rowspan="3">x<td>y<tr><tr><td>z</table>
"""


def _merged_xlsx(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Merged"
    sheet.append(["Region", "Item", "Value"])
    for row in range(8):
        sheet.append([f"r{row}", f"item {row}", row * 10])
    sheet.merge_cells("A2:C4")  # rows 3 and 4 have no cells of their own
    sheet.merge_cells("A6:A7")
    sheet.merge_cells("B8:C9")  # the last row holds only the first column's cell
    book.save(path)
    return path


def _docx_cell(text: str, merge: str = "") -> str:
    merge = {"": "", "start": '<w:vMerge w:val="restart"/>', "continue": "<w:vMerge/>"}[merge]
    paragraph = f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" if text else "<w:p/>"
    return f"<w:tc><w:tcPr>{merge}</w:tcPr>{paragraph}</w:tc>"


MERGED_DOCX_BODY = (
    "<w:tbl><w:tr><w:trPr><w:tblHeader/></w:trPr>"
    + _docx_cell("First")
    + _docx_cell("Second")
    + "</w:tr><w:tr>"
    + _docx_cell("merged down", "start")
    + _docx_cell("also merged", "start")
    + "</w:tr><w:tr>"
    + _docx_cell("", "continue")
    + _docx_cell("", "continue")
    + "</w:tr><w:tr>"
    + _docx_cell("after")
    + _docx_cell("the merge")
    + "</w:tr></w:tbl>"
)


# Tables Chromium tags or leaves out on its own judgement, among the ones
# needing their spans written as laid out: a table left out of the tagged PDF
# must not change how the next ones are tagged. Each later table's rows fill
# the same columns, so any grid problem is one the export made.
MERGED_PROBES = """# Merged cells among other tables

<table><tr><td>Note: a boxed callout</td></tr></table>

<table hidden><tr><td rowspan="2">hidden</td></tr><tr></tr></table>

<table aria-hidden="true"><tr><td>not for assistive technology</td></tr></table>

<table role="presentation"><tr><td>for layout</td></tr></table>

<details><summary>More</summary><table><tr><td>closed</td></tr></table></details>

<table></table>

<table><tr><td>one</td></tr><tr><td>column</td></tr></table>

<table><tbody><tr><td rowspan="3">cut at its body</td><td>c1</td></tr></tbody>
<tbody><tr><td>c2</td><td>c3</td></tr><tr><td>c4</td><td>c5</td></tr></tbody></table>

<table><tbody><tr><td rowspan="3">three rows</td><td>d1</td></tr><tr><td>d2</td></tr>
<tr><td>d3</td></tr></tbody></table>

<table><thead><tr><th>A</th><th>B</th></tr></thead><tbody><tr><td>a</td><td>b</td></tr>
<tr><td colspan="2" rowspan="3">the only cell in its row</td></tr><tr></tr><tr></tr>
<tr><td>c</td><td>d</td></tr></tbody></table>

<table><tfoot><tr><td rowspan="2">foot</td><td>f</td></tr></tfoot>
<thead><tr><th>H</th><th>I</th></tr></thead>
<tbody><tr><td rowspan="3">body</td><td>1</td></tr><tr></tr><tr><td>2</td></tr></tbody>
<thead><tr><th rowspan="2">second head</th><th>J</th></tr></thead>
<tfoot><tr><td rowspan="2">second foot</td><td>x</td></tr><tr><td>y</td></tr></tfoot>
<tbody><tr><td>7</td><td>8</td></tr></tbody></table>
"""


def page_images(path: Path) -> list[str]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        return [
            hashlib.sha256(document[i].render(scale=2).to_pil().tobytes()).hexdigest()
            for i in range(len(document))
        ]
    finally:
        document.close()


def _merged_source(kind: str, root: Path) -> Path:
    if kind in ("markdown", "probes"):
        source = root / f"{kind}.md"
        source.write_text(
            MERGED_MARKDOWN if kind == "markdown" else MERGED_PROBES, encoding="utf-8"
        )
        return source
    if kind == "xlsx":
        return _merged_xlsx(root / "merged.xlsx")
    return make_docx(root / "merged.docx", body=MERGED_DOCX_BODY)


@pytest.mark.integration
@pytest.mark.parametrize("kind", ["markdown", "probes", "xlsx", "docx"])
@pytest.mark.usefixtures("streaming_finish")
def test_merged_cells_are_tagged_with_the_rows_the_table_holds(tmp_path, monkeypatch, kind):
    source = _merged_source(kind, tmp_path)
    settings = make_settings(document_language="en")
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        tables = table_cells(reader)
    spans = [down for rows in tables for row in rows for down, _ in row if down != 1]
    # The document's merge covers whole rows, so its cells span one tagged row.
    assert bool(spans) == (kind != "docx")
    if kind == "probes":
        assert [[(3, 1), (1, 1)], [(1, 1)], [(1, 1)]] in tables  # its span kept whole
    for rows in tables:
        assert grid_problems(rows) == [], rows
    _keep_evidence(final, f"fixture-merged-{kind}")

    # The pages are the ones the markup as written prints.
    monkeypatch.setattr(table_spans, "normalize", lambda html, *args, **kwargs: html)
    result = export(
        ExportRequest(str(source), str(tmp_path / "as-written"), settings, browser="bundled")
    )
    assert result.exit_code == 0, result.report
    as_written = next((tmp_path / "as-written").glob("*.pdf"))
    assert page_images(final) == page_images(as_written)


# A Word table whose rows do not fill one grid. The last row continues each
# column's merge one column wide, so Mammoth extends both the wide merged cell
# and the merged cell beside it, a row further up. The wide cell used to span
# into the column that cell holds, and the two overlapped.
IRREGULAR_DOCX_BODY = merge_table(
    ["Name", "Value"],
    ["alpha", ("beside the merge", 1, "r")],
    [("merged across both columns", 2, "r")],
    [("", 1, "c"), ("", 1, "c")],
    header=1,
)


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
def test_an_irregular_docx_table_is_tagged_as_one_grid(tmp_path):
    source = make_docx(tmp_path / "irregular.docx", body=IRREGULAR_DOCX_BODY)
    settings = make_settings(document_language="en")
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    (record,) = result.report["documents"]
    assert any(warning.startswith("DOCX_TABLE_OVERLAP: Table 1 ") for warning in record["warnings"])
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        (rows,) = table_cells(reader)
        text = "".join(page.extract_text() or "" for page in reader.pages)
    assert grid_problems(rows) == [], rows
    assert rows == [[(1, 1), (1, 1)], [(1, 1), (2, 1)], [(1, 1)]]
    for words in ("alpha", "beside the merge", "merged across both columns"):
        assert words in " ".join(text.split())
    _keep_evidence(final, "fixture-docx-irregular")


def _outer_table(inner: str, row: int) -> str:
    """A five-row table of two columns whose second cell in ``row`` holds ``inner``."""
    cells = [
        [
            f"<w:tc><w:tcPr/>{inner}<w:p/></w:tc>"
            if (r, c) == (row, 1)
            else f"<w:tc><w:tcPr/><w:p><w:r><w:t>o{r}{c}</w:t></w:r></w:p></w:tc>"
            for c in range(2)
        ]
        for r in range(5)
    ]
    return "<w:tbl>%s</w:tbl>" % "".join(f"<w:tr>{''.join(row)}</w:tr>" for row in cells)


_INNER_SHORT = merge_table(
    ["n1", "n2", "n3"], ["n4", "n5", "n6"], ["n7"], ["n8", "n9", "n10"], ["n11", "n12"]
)

# Word tables with short rows, which Word draws short and leaves blank beyond
# their last cell: a striped body row, a header row, a row longer than the
# rest, a row that stops short beside a merged cell, a table nested in a
# striped outer row and in a plain one, and a table wide enough for a plate.
SHORT_ROWS_DOCX_BODY = "<w:p/>".join(
    [
        merge_table(
            ["Name", "Value", "Note"],
            ["alpha", "1", "first"],
            ["beta", "2"],
            ["gamma", "3", "third"],
            ["delta"],
            header=1,
        ),
        merge_table(["a", "b"], ["c", "d"], ["e", "f", "g", "h"], ["i", "j"]),
        merge_table(["a", "b", "c", "d"], [("e", 1, "r"), "f"], [("", 1, "c"), "g", "h", "i"]),
        _outer_table(_INNER_SHORT, 1),
        _outer_table(_INNER_SHORT, 2),
        merge_table(
            [f"h{c}" for c in range(16)],
            [f"w{c}" for c in range(16)],
            [f"s{c}" for c in range(9)],
            [f"w{c}" for c in range(16)],
            [f"s{c}" for c in range(5)],
            header=1,
        ),
    ]
)
# A header row shorter than the body, after the plate so that it prints on a
# page of its own. Filling it out changes the page: with two header cells the
# table was taken for a two-column one and its cells given a minimum width.
SHORT_HEADER_DOCX_BODY = merge_table(
    ["Heading one", "Heading two"],
    ["alpha", "beta", "gamma"],
    ["delta", "epsilon", "zeta"],
    header=1,
)
_BUNDLED_THEMES = sorted(name for name, theme in discover_themes().items() if is_builtin(theme))


@pytest.mark.integration
@pytest.mark.parametrize("theme", [*_BUNDLED_THEMES, "tinted", "source"])
@pytest.mark.usefixtures("streaming_finish")
@pytest.mark.filterwarnings("ignore::platefold.theme.ThemeContrastWarning")
def test_docx_short_rows_are_tagged_full_and_print_as_word_draws_them(tmp_path, monkeypatch, theme):
    # A custom theme on tinted paper, and the Word file's own palette.
    if theme == "tinted":
        tinted = json.loads((BUILTIN_DIR / "default.json").read_text(encoding="utf-8"))
        tinted.update(name="tinted-proof", label="Tinted proof")
        tinted["colors"].update(page="#e9f1dc", surface="#d5e3c0", surface_alt="#dde9cb")
        (tmp_path / "tinted.json").write_text(json.dumps(tinted), encoding="utf-8")
        options = {"theme": str(tmp_path / "tinted.json")}
    else:
        options = {"source_theme": True} if theme == "source" else {"theme": theme}
    body = SHORT_ROWS_DOCX_BODY + "<w:p/>" + SHORT_HEADER_DOCX_BODY
    source = make_docx(tmp_path / "short.docx", body=body)
    settings = make_settings(document_language="en", **options)
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    (record,) = result.report["documents"]
    assert not any(warning.startswith("DOCX_") for warning in record["warnings"])
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        tables = table_cells(reader)
        header = _first_row_roles(reader)[-1]
    assert len(tables) == 9  # the nested tables are tagged tables of their own
    for rows in tables:
        assert grid_problems(rows) == [], rows
    # The short header row keeps its two header cells; the one that fills it is data.
    assert header == ["/TH", "/TH", "/TD"]
    if theme == "default":
        _keep_evidence(final, "fixture-docx-short-rows")

    # The cells that fill the short rows draw nothing: the pages are the ones
    # the rows print as written, with no border or stripe where Word has none,
    # and whatever lies behind a table, page or outer stripe, showing as before.
    # The last page holds the short header row, where that is not so.
    monkeypatch.setattr(docx._Grid, "padding", lambda self, budget: [])
    result = export(
        ExportRequest(str(source), str(tmp_path / "short"), settings, browser="bundled")
    )
    assert result.exit_code == 0, result.report
    as_written = next((tmp_path / "short").glob("*.pdf"))
    with PdfReader(as_written) as reader:
        assert any(grid_problems(rows) for rows in table_cells(reader))
    padded, unpadded = page_images(final), page_images(as_written)
    assert len(padded) == len(unpadded)
    assert padded[:-1] == unpadded[:-1]

    # With two header cells, the short header made the table look like a
    # two-column one, and its cells were given a minimum width. Filled out,
    # its columns are sized as for the same table with a full header row.
    if theme == "default":
        monkeypatch.undo()
        full = SHORT_HEADER_DOCX_BODY.replace(
            "Heading two</w:t></w:r></w:p></w:tc>",
            "Heading two</w:t></w:r></w:p></w:tc><w:tc><w:tcPr/><w:p/></w:tc>",
            1,
        )
        reference = make_docx(tmp_path / "full.docx", body=SHORT_ROWS_DOCX_BODY + "<w:p/>" + full)
        result = export(
            ExportRequest(str(reference), str(tmp_path / "full"), settings, browser="bundled")
        )
        assert result.exit_code == 0, result.report
        columns = [
            _word_columns(path, ["alpha", "beta", "gamma"])
            for path in (final, next((tmp_path / "full").glob("*.pdf")), as_written)
        ]
        assert columns[0] == columns[1] != columns[2]


def _first_row_roles(reader: PdfReader) -> list[list[str]]:
    """For each tagged table, its first row's cell roles."""
    roles, pending = [], [reader.root_object["/StructTreeRoot"].get_object()]
    while pending:
        node = pending.pop(0)
        if node.get("/S") == "/Table":
            groups = _elements(node)
            while groups and groups[0].get("/S") != "/TR":
                groups[:1] = _elements(groups[0])
            if groups:
                roles.append(
                    [str(c["/S"]) for c in _elements(groups[0]) if c.get("/S") in ("/TD", "/TH")]
                )
        pending[:0] = _elements(node)
    return roles


def _word_columns(path: Path, words: list[str]) -> list[float]:
    """Where each of ``words`` starts on the last page, in points from the left."""
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        found = {word["text"]: round(word["x0"], 1) for word in pdf.pages[-1].extract_words()}
    return [found[word] for word in words]


_BORDERS = "".join(
    f'<w:{edge} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
)
# A staggered Word form: rows that start one or two columns in (w:gridBefore),
# rows that stop early (w:gridAfter), and a merge down an indented column. The
# grid and borders are there for Word and LibreOffice; Plate|Fold sizes its own.
INDENTED_DOCX_BODY = (
    f"<w:tbl><w:tblPr><w:tblBorders>{_BORDERS}</w:tblBorders></w:tblPr><w:tblGrid>"
    + '<w:gridCol w:w="2200"/>' * 4
    + "</w:tblGrid>"
    + indented_row("Field", "Value", "Notes", "Owner", header=True)
    + indented_row("Subitem", "twelve", "extra", before=1)
    + indented_row("Detail", "why", before=2)
    + indented_row("alpha", "beta", "gamma", "delta")
    + indented_row("Total", "fortytwo", after=2)
    + indented_row(("Merged", 1, "r"), "left", "right", before=1)
    + indented_row(("", 1, "c"), "lower", "last", before=1)
    + "</w:tbl>"
)
# Each word of the form and the grid column Word puts it in, counted from 0.
INDENTED_COLUMNS = {
    "Field": 0, "Value": 1, "Notes": 2, "Owner": 3,
    "Subitem": 1, "twelve": 2, "extra": 3,
    "Detail": 2, "why": 3,
    "alpha": 0, "beta": 1, "gamma": 2, "delta": 3,
    "Total": 0, "fortytwo": 1,
    "Merged": 1, "left": 2, "right": 3,
    "lower": 2, "last": 3,
}  # fmt: skip


def word_columns(path: Path) -> dict[str, int]:
    """The column each word of the indented form starts in, judged by the header words."""
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        words = {w["text"]: w["x0"] for page in pdf.pages for w in page.extract_words()}
    starts = sorted(words[name] for name in ("Field", "Value", "Notes", "Owner"))
    return {
        word: max(i for i, start in enumerate(starts) if words[word] >= start - 1)
        for word in INDENTED_COLUMNS
    }


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
def test_indented_docx_rows_print_and_tag_in_their_own_columns(tmp_path):
    source = make_docx(tmp_path / "indented.docx", body=INDENTED_DOCX_BODY)
    settings = make_settings(document_language="en")
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    (record,) = result.report["documents"]
    assert not any(warning.startswith("DOCX_") for warning in record["warnings"])
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        (rows,) = table_cells(reader)
        (header,) = _first_row_roles(reader)
    assert grid_problems(rows) == [], rows
    assert header == ["/TH", "/TH", "/TH", "/TH"]
    # The merge down the indented column spans its two rows.
    assert rows[5] == [(1, 1), (2, 1), (1, 1), (1, 1)]
    assert word_columns(final) == INDENTED_COLUMNS
    _keep_evidence(final, "fixture-docx-indented")


def _last_line_clearance(path: Path) -> float:
    """The least room, in points, between a merged cell's last line and the rule below it."""
    import pdfplumber

    least = float("inf")
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            words = page.extract_words()
            rules = sorted(r["top"] for r in page.rects if r["width"] > 200 and r["height"] < 2)
            numbers = [w for w in words if w["text"].startswith("n") and w["text"][1:].isdigit()]
            if numbers:
                last = max(numbers, key=lambda word: word["top"])
                below = [top for top in rules if top > last["top"]]
                if below:
                    least = min(least, below[0] - last["bottom"])
    return least


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
def test_a_merged_cell_broken_across_pages_keeps_its_last_line_clear(tmp_path):
    # Chromium drew the row after such a cell over its last lines when the
    # rows it spanned had no cells; spanning the rows laid out, it does not.
    rows = "".join(f"<tr><td>x{i}</td><td>b</td><td>c</td></tr>" for i in range(30))
    head = "<thead><tr><th>A</th><th>B</th><th>C</th></tr></thead>"
    for count in range(60, 80, 2):
        tall = "<br>".join(f"n{i}" for i in range(count))
        block = f'<tr><td rowspan="3" colspan="3">{tall}</td></tr><tr></tr><tr></tr>'
        source = tmp_path / f"tall-{count}.md"
        source.write_text(f"# T\n\n<table>{head}<tbody>{rows}{block}{rows}</tbody></table>\n")
        out = tmp_path / f"out-{count}"
        settings = make_settings(document_language="en")
        assert (
            export(ExportRequest(str(source), str(out), settings, browser="bundled")).exit_code == 0
        )
        assert _last_line_clearance(next(out.glob("*.pdf"))) > 0, count


@pytest.mark.parametrize(
    "html,expected",
    [
        ("<td></td>", "<td>&#160;</td>"),
        ("<td class='x'> \n </td>", "<td class='x'> \n &#160;</td>"),
        ('<th scope="col"></th>', '<th scope="col">&#160;</th>'),
        ("<td><p></p></td>", "<td><p></p>&#160;</td>"),
        ("<td><br></td>", "<td><br>&#160;</td>"),
        ("<td>text</td>", "<td>text</td>"),
        ("<td><b>x</b></td>", "<td><b>x</b></td>"),
        ("<td>&nbsp;</td>", "<td>&nbsp;</td>"),
        ("<td><img src='a.png' alt=''></td>", "<td><img src='a.png' alt=''></td>"),
        ("<td><svg></svg></td>", "<td><svg></svg></td>"),
        ("<tr><td></td><td>b</td></tr>", "<tr><td>&#160;</td><td>b</td></tr>"),
        ("<thead><tr><th></th></tr></thead>", "<thead><tr><th>&#160;</th></tr></thead>"),
    ],
)
def test_empty_cells_get_a_no_break_space(html, expected):
    assert fill_empty_cells(html) == expected


def test_empty_cell_filling_is_linear_on_hostile_input():
    started = time.perf_counter()
    fill_empty_cells("<td>" + " " * 200_000 + "x</td>" + "<td>" + "<i>" * 50_000 + "y</td>")
    assert time.perf_counter() - started < 2


def test_non_standard_roles_map_to_standard_types():
    mapping = role_map({"/Strong", "/Em", "/P", "/Invented"}, {"/Custom": "/Span"})
    assert mapping == {
        "/Custom": "/Span",
        "/Em": "/Span",
        "/Strong": "/Span",
        "/Invented": "/NonStruct",
    }
    assert set(mapping.values()) <= STANDARD_TYPES


# Raw HTML tables in Markdown with rows of different widths: short and long
# body rows, a gap beside a cell spanning down, a short row striped as an even
# row, and a short nested table. The header is full: a short one also changes
# how the columns are sized, as test_docx_short_rows_are_tagged_full shows.
SHORT_ROWS_MARKDOWN = """# Short rows
Language: en

<table><thead><tr><th>Field</th><th>Value</th><th>Note</th><th>Owner</th></tr></thead><tbody>
<tr><td>alpha</td><td>beta</td><td>gamma</td></tr>
<tr><td>short</td></tr>
<tr><td>a</td><td>b</td><td rowspan="2">spans</td><td>d</td></tr>
<tr><td>e</td></tr>
<tr><td>outer</td><td><table><tr><td>n1</td><td>n2</td></tr><tr><td>n3</td></tr></table></td></tr>
</tbody></table>

After the table.
"""


@pytest.mark.integration
@pytest.mark.usefixtures("streaming_finish")
def test_markdown_short_rows_are_tagged_full_and_print_as_written(tmp_path, monkeypatch):
    source = tmp_path / "short.md"
    source.write_text(SHORT_ROWS_MARKDOWN, encoding="utf-8")
    settings = make_settings(document_language="en")
    result = export(ExportRequest(str(source), str(tmp_path / "out"), settings, browser="bundled"))
    assert result.exit_code == 0, result.report
    final = next((tmp_path / "out").glob("*.pdf"))
    with PdfReader(final) as reader:
        tables = table_cells(reader)
    assert len(tables) == 2  # the nested table is tagged on its own
    for rows in tables:
        assert grid_problems(rows) == [], rows
    _keep_evidence(final, "fixture-markdown-short-rows")

    # The cells that fill the short rows draw nothing: the pages are the ones
    # the rows print as written.
    monkeypatch.setattr(table_spans, "_short_rows", lambda table, budget: [])
    result = export(
        ExportRequest(str(source), str(tmp_path / "short"), settings, browser="bundled")
    )
    assert result.exit_code == 0, result.report
    as_written = next((tmp_path / "short").glob("*.pdf"))
    with PdfReader(as_written) as reader:
        assert any(grid_problems(rows) for rows in table_cells(reader))
    assert page_images(final) == page_images(as_written)
